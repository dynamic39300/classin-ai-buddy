#!/usr/bin/env python3
"""Profile a ClassIn IM XLSX without emitting raw messages or identifiers.

Created during Round 0. The read-only structural scan may be reused only with
formal input fingerprinting and audit logging; see tools/analysis/README.md.
"""

from __future__ import annotations

import argparse
import json
import re
import time
from collections import Counter, defaultdict
from datetime import date, datetime
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


SAFE_CATEGORICAL_FIELDS = (
    "identity",
    "user_type",
    "user_num",
    "clustertype",
    "msgcmd",
    "dt",
    "rn",
)


def field_key(header: Any) -> str:
    text = "" if header is None else str(header).strip()
    return re.split(r"[（(]", text, maxsplit=1)[0].strip().lower()


def json_value(value: Any) -> Any:
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    return value


def shape_signature(value: Any) -> str:
    """Return a non-reversible structural signature for a sensitive value."""
    if value is None or value == "":
        return "EMPTY"
    text = str(value)
    text = re.sub(r"[A-Za-z]", "A", text)
    text = re.sub(r"\d", "9", text)
    text = re.sub(r"[\u4e00-\u9fff]", "中", text)
    text = re.sub(r"(.)\1{3,}", r"\1×N", text)
    return text[:80]


def content_shape(value: Any) -> str:
    if value is None or value == "":
        return "EMPTY"
    text = str(value).strip()
    if not text:
        return "BLANK"
    lowered = text.lower()
    if lowered.startswith(("http://", "https://")):
        return "URL"
    if text.startswith("{") and text.endswith("}"):
        return "JSON_OBJECT"
    if text.startswith("[") and text.endswith("]"):
        return "JSON_ARRAY"
    if text.startswith("<") and text.endswith(">"):
        return "MARKUP"
    has_cjk = bool(re.search(r"[\u4e00-\u9fff]", text))
    has_latin = bool(re.search(r"[A-Za-z]", text))
    has_digit = bool(re.search(r"\d", text))
    if has_cjk and has_latin:
        return "CJK_LATIN_MIXED"
    if has_cjk:
        return "CJK_TEXT"
    if has_latin:
        return "LATIN_TEXT"
    if has_digit:
        return "NUMERIC_OR_SYMBOL"
    return "SYMBOL_OR_EMOJI"


def length_bucket(length: int) -> str:
    if length == 0:
        return "0"
    if length <= 5:
        return "1-5"
    if length <= 20:
        return "6-20"
    if length <= 50:
        return "21-50"
    if length <= 100:
        return "51-100"
    if length <= 300:
        return "101-300"
    return "301+"


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("input_xlsx", type=Path)
    parser.add_argument("--output", type=Path)
    args = parser.parse_args()

    started = time.time()
    workbook = load_workbook(args.input_xlsx, read_only=True, data_only=True)
    worksheet = workbook.active
    rows = worksheet.iter_rows(values_only=True)
    raw_headers = next(rows)
    keys = [field_key(value) for value in raw_headers]
    index = {key: position for position, key in enumerate(keys)}

    counts: dict[str, Counter[str]] = {
        field: Counter() for field in SAFE_CATEGORICAL_FIELDS if field in index
    }
    shape_counts: dict[str, Counter[str]] = {
        "targetuids": Counter(),
        "msgdata": Counter(),
        "concent": Counter(),
    }
    length_counts: dict[str, Counter[str]] = {
        "msgdata": Counter(),
        "concent": Counter(),
    }
    nonempty = Counter()
    unique_values: dict[str, set[Any]] = defaultdict(set)
    unique_fields = ("clusterid", "sourceuid", "id", "msgid")
    rows_by_type = Counter()
    cluster_message_counts = Counter()
    cluster_type_by_id: dict[Any, str] = {}
    cluster_first_rank: dict[Any, tuple[int, str]] = {}
    cluster_last_rank: dict[Any, tuple[int, str]] = {}
    rank_sum_by_date = Counter()
    rank_count_by_date = Counter()

    total_rows = 0
    for row in rows:
        total_rows += 1
        for field, counter in counts.items():
            value = row[index[field]]
            counter[str(json_value(value)) if value not in (None, "") else "<EMPTY>"] += 1

        for field in unique_fields:
            if field in index:
                value = row[index[field]]
                if value not in (None, ""):
                    unique_values[field].add(value)

        if "clustertype" in index:
            value = row[index["clustertype"]]
            rows_by_type[str(value) if value not in (None, "") else "<EMPTY>"] += 1

        cluster_id = row[index["clusterid"]] if "clusterid" in index else None
        cluster_type = row[index["clustertype"]] if "clustertype" in index else None
        if cluster_id not in (None, ""):
            cluster_message_counts[cluster_id] += 1
            cluster_type_by_id[cluster_id] = str(cluster_type)

        rank_value = row[index["rn"]] if "rn" in index else None
        date_value = row[index["dt"]] if "dt" in index else None
        try:
            rank_number = int(rank_value)
        except (TypeError, ValueError):
            rank_number = None
        date_text = str(json_value(date_value)) if date_value not in (None, "") else "<EMPTY>"
        if rank_number is not None:
            rank_sum_by_date[date_text] += rank_number
            rank_count_by_date[date_text] += 1
            if cluster_id not in (None, ""):
                current_first = cluster_first_rank.get(cluster_id)
                current_last = cluster_last_rank.get(cluster_id)
                if current_first is None or rank_number < current_first[0]:
                    cluster_first_rank[cluster_id] = (rank_number, date_text)
                if current_last is None or rank_number > current_last[0]:
                    cluster_last_rank[cluster_id] = (rank_number, date_text)

        for field in ("targetuids", "msgdata", "concent"):
            if field not in index:
                continue
            value = row[index[field]]
            if value not in (None, ""):
                nonempty[field] += 1
            if field == "targetuids":
                shape_counts[field][shape_signature(value)] += 1
            else:
                shape_counts[field][content_shape(value)] += 1
                length_counts[field][length_bucket(len(str(value)) if value is not None else 0)] += 1

    result = {
        "source": {
            "file_name": args.input_xlsx.name,
            "file_size_bytes": args.input_xlsx.stat().st_size,
            "sheet": worksheet.title,
            "declared_max_row": worksheet.max_row,
            "declared_max_column": worksheet.max_column,
            "data_rows_scanned": total_rows,
            "elapsed_seconds": round(time.time() - started, 2),
        },
        "schema": [
            {"position": position + 1, "key": key, "header": str(raw_headers[position])}
            for position, key in enumerate(keys)
        ],
        "categorical_counts": {
            field: dict(counter.most_common()) for field, counter in counts.items()
        },
        "unique_counts": {field: len(values) for field, values in unique_values.items()},
        "sampling_structure": {
            "cluster_size_distribution": dict(Counter(cluster_message_counts.values()).most_common()),
            "clusters_by_type": dict(Counter(cluster_type_by_id.values()).most_common()),
            "first_rank_distribution": dict(Counter(value[0] for value in cluster_first_rank.values()).most_common()),
            "last_rank_distribution": dict(Counter(value[0] for value in cluster_last_rank.values()).most_common()),
            "first_rank_date_distribution": dict(Counter(value[1] for value in cluster_first_rank.values()).most_common()),
            "last_rank_date_distribution": dict(Counter(value[1] for value in cluster_last_rank.values()).most_common()),
            "mean_rank_by_date": {
                day: round(rank_sum_by_date[day] / count, 2)
                for day, count in sorted(rank_count_by_date.items())
                if count
            },
        },
        "nonempty_counts": dict(nonempty),
        "shape_counts": {
            field: dict(counter.most_common(30)) for field, counter in shape_counts.items()
        },
        "length_buckets": {
            field: dict(counter.most_common()) for field, counter in length_counts.items()
        },
        "privacy_note": "No raw message, user identifier, cluster identifier, or target identifier is emitted.",
    }

    serialized = json.dumps(result, ensure_ascii=False, indent=2)
    if args.output:
        args.output.parent.mkdir(parents=True, exist_ok=True)
        args.output.write_text(serialized, encoding="utf-8")
    else:
        print(serialized)


if __name__ == "__main__":
    main()
