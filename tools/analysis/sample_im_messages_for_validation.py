#!/usr/bin/env python3
"""Create a Round 0 redacted sample for the archived rule-based IM taxonomy.

This is a discovery sample, not a representative or decision-grade sample. See
tools/analysis/README.md before use.
"""

from __future__ import annotations

import argparse
import json
import random
import re
from collections import defaultdict
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from analyze_im_chat_xlsx import TOPIC_PATTERNS, channel_name, content_shape, field_key, safe_normalize


SURNAME_CLASS = "赵钱孙李周吴郑王冯陈褚卫蒋沈韩杨朱秦尤许何吕施张孔曹严华金魏陶姜戚谢邹喻柏水窦章云苏潘葛奚范彭郎鲁韦昌马苗凤花方俞任袁柳鲍史唐费廉岑薛雷贺倪汤滕殷罗毕郝邬安常乐于时傅皮卞齐康伍余元卜顾孟平黄和穆萧尹姚邵湛汪祁毛禹狄米贝明臧计伏成戴谈宋茅庞熊纪舒屈项祝董梁杜阮蓝闵季贾路江童颜郭梅盛林刁钟徐邱骆高夏蔡田樊胡凌霍虞万支柯管卢莫经房裘缪干解应宗丁宣贲邓郁单杭洪包诸左石崔吉钮龚程嵇邢裴陆荣翁荀羊甄家封芮储靳邴松井段富巫乌焦巴弓牧隗山谷车侯宓蓬全郗班仰秋仲伊宫宁仇栾暴甘钭厉戎祖武符刘景詹束龙叶幸司韶郜黎蓟薄印宿白怀蒲台从鄂索咸籍赖卓蔺屠蒙池乔阴胥能苍双闻莘党翟谭贡劳逄姬申扶堵冉宰郦雍郤璩桑桂濮牛寿通边扈燕冀郏浦尚农温别庄晏柴瞿阎充慕连茹习宦艾鱼容向古易慎戈廖庾终暨居衡步都耿满弘匡国文寇广禄阙东欧殳沃利蔚越夔隆师巩厍聂晁勾敖融冷訾辛阚那简饶空曾毋沙乜养鞠须丰巢关蒯相查后荆红游竺权逯盖益桓公"
NAME_PATTERN = re.compile(rf"(?<![\u4e00-\u9fff])[{SURNAME_CLASS}][\u4e00-\u9fff]{{1,2}}(?:老师|同学)?(?![\u4e00-\u9fff])")


def redact(text: str) -> str:
    text = safe_normalize(text)
    text = NAME_PATTERN.sub("<NAME>", text)
    text = re.sub(r"\b[a-z][a-z0-9._-]{5,}\b", "<TOKEN>", text, flags=re.IGNORECASE)
    return text[:180]


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("input_xlsx", type=Path)
    parser.add_argument("--output", required=True, type=Path)
    parser.add_argument("--per-cell", type=int, default=5)
    args = parser.parse_args()

    rng = random.Random(20260828)
    workbook = load_workbook(args.input_xlsx, read_only=True, data_only=True)
    worksheet = workbook.active
    rows = worksheet.iter_rows(values_only=True)
    headers = next(rows)
    index = {field_key(value): position for position, value in enumerate(headers)}

    samples: dict[tuple[str, str], list[str]] = defaultdict(list)
    seen: dict[tuple[str, str], int] = defaultdict(int)
    for row in rows:
        channel = channel_name(row[index["clustertype"]])
        value = row[index["concent"]]
        text = "" if value is None else str(value)
        if content_shape(value) in {"空值", "空白", "JSON对象", "结构化数组", "标记文本", "符号或表情"}:
            continue
        normalized = safe_normalize(text)
        for topic, pattern in TOPIC_PATTERNS:
            if not pattern.search(normalized):
                continue
            key = (channel, topic)
            seen[key] += 1
            item = redact(text)
            bucket = samples[key]
            if len(bucket) < args.per_cell:
                bucket.append(item)
            else:
                position = rng.randrange(seen[key])
                if position < args.per_cell:
                    bucket[position] = item

    result: dict[str, Any] = defaultdict(dict)
    for (channel, topic), values in samples.items():
        result[channel][topic] = {"classified_count": seen[(channel, topic)], "redacted_samples": values}
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8")


if __name__ == "__main__":
    main()
