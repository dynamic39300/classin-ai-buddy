#!/usr/bin/env python3
"""Extract Pilot-0 review rows with participant pseudonyms and direct-ID redaction.

The output remains restricted research material because free text can contain
contextual personal information even after automatic redaction. Never commit the
output JSON or derived review workbook to the repository.
"""

from __future__ import annotations

import argparse
import hashlib
import json
import re
import sys
import time
from collections import Counter, defaultdict
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


EXPECTED_SHA256 = "a28f4c3125326c2e0f9086f8c0f67104671b94b2992a0173f4c34f90690fee6c"

PHONE_PATTERN = re.compile(r"(?<!\d)1[3-9]\d{9}(?!\d)")
ID_CARD_PATTERN = re.compile(r"(?<!\d)\d{17}[\dXx](?!\d)")
EMAIL_PATTERN = re.compile(r"[A-Za-z0-9._%+-]+@[A-Za-z0-9.-]+\.[A-Za-z]{2,}")
URL_PATTERN = re.compile(r"https?://[^\s<>\]\[\"']+", re.IGNORECASE)
ACCOUNT_PATTERN = re.compile(
    r"(?i)(微信|wechat|qq|钉钉|手机号|电话|联系方式)\s*[:：]?\s*[A-Za-z0-9_.-]{4,}"
)
HONORIFIC_NAME_PATTERN = re.compile(
    r"(?<![\u4e00-\u9fff])([\u4e00-\u9fff]{2,4})(老师|同学|家长)(?![\u4e00-\u9fff])"
)


def file_sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def field_key(header: Any) -> str:
    text = "" if header is None else str(header).strip()
    return re.split(r"[（(]", text, maxsplit=1)[0].strip().lower()


def is_blank(value: Any) -> bool:
    return value is None or value == "" or (isinstance(value, str) and not value.strip())


def parse_integer(value: Any) -> int | None:
    if value is None or value == "" or isinstance(value, bool):
        return None
    if isinstance(value, int):
        return value
    if isinstance(value, float) and value.is_integer():
        return int(value)
    text = str(value).strip()
    return int(text) if re.fullmatch(r"[-+]?\d+", text) else None


def excel_column_name(index: int) -> str:
    result = ""
    value = index
    while value:
        value, remainder = divmod(value - 1, 26)
        result = chr(65 + remainder) + result
    return result


def pseudonym(index: int) -> str:
    return f"参与者{excel_column_name(index)}"


def canonical_content(msgdata: Any, content: Any) -> tuple[str, str]:
    if isinstance(msgdata, str):
        stripped = msgdata.strip()
        if stripped.startswith("{") and stripped.endswith("}"):
            try:
                parsed = json.loads(stripped)
            except (json.JSONDecodeError, TypeError, ValueError):
                parsed = None
            if isinstance(parsed, dict) and isinstance(parsed.get("content"), str):
                return parsed["content"], "msgdata.content"
    if not is_blank(content):
        return str(content), "concent_fallback"
    if not is_blank(msgdata):
        return str(msgdata), "msgdata_raw_fallback"
    return "[无可用正文]", "no_content"


def redact_text(
    text: str,
    replacements: dict[str, str],
) -> tuple[str, dict[str, int]]:
    redacted = text
    counts = Counter()
    for original, replacement in sorted(
        replacements.items(), key=lambda item: len(item[0]), reverse=True
    ):
        if not original or len(original) < 2:
            continue
        occurrences = redacted.count(original)
        if occurrences:
            redacted = redacted.replace(original, replacement)
            counts["participant_reference"] += occurrences

    for name, pattern, replacement in (
        ("url", URL_PATTERN, "[链接]"),
        ("email", EMAIL_PATTERN, "[邮箱]"),
        ("id_card", ID_CARD_PATTERN, "[证件号]"),
        ("phone", PHONE_PATTERN, "[手机号]"),
        ("account", ACCOUNT_PATTERN, "[账号]"),
    ):
        redacted, replacements_count = pattern.subn(replacement, redacted)
        counts[name] += replacements_count

    def replace_honorific(match: re.Match[str]) -> str:
        counts["honorific_name"] += 1
        return f"[姓名]{match.group(2)}"

    redacted = HONORIFIC_NAME_PATTERN.sub(replace_honorific, redacted)
    return redacted, dict(counts)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("input_xlsx", type=Path)
    parser.add_argument("restricted_map", type=Path)
    parser.add_argument("public_manifest", type=Path)
    parser.add_argument("--output", type=Path, required=True)
    parser.add_argument("--expected-sha256", default=EXPECTED_SHA256)
    parser.add_argument("--progress-every", type=int, default=100_000)
    args = parser.parse_args()

    started = time.time()
    actual_sha256 = file_sha256(args.input_xlsx)
    if actual_sha256 != args.expected_sha256:
        raise SystemExit("Input fingerprint mismatch")

    restricted_map = json.loads(args.restricted_map.read_text(encoding="utf-8"))
    public_manifest = json.loads(args.public_manifest.read_text(encoding="utf-8"))
    if restricted_map["source_sha256"] != actual_sha256:
        raise SystemExit("Restricted map source fingerprint mismatch")
    if public_manifest["source"]["sha256"] != actual_sha256:
        raise SystemExit("Public manifest source fingerprint mismatch")

    restricted_by_cluster = {
        record["clusterid"]: record for record in restricted_map["samples"]
    }
    public_by_sample = {
        record["sample_id"]: record for record in public_manifest["samples"]
    }
    raw_rows_by_sample: dict[str, list[dict[str, Any]]] = defaultdict(list)

    workbook = load_workbook(args.input_xlsx, read_only=True, data_only=False)
    worksheet = workbook.worksheets[0]
    rows = worksheet.iter_rows(values_only=True)
    headers = next(rows)
    keys = [field_key(value) for value in headers]
    index = {key: position for position, key in enumerate(keys)}

    scanned_rows = 0
    for row in rows:
        scanned_rows += 1
        cluster_id = row[index["clusterid"]]
        selection = restricted_by_cluster.get(cluster_id)
        if selection is None:
            if args.progress_every and scanned_rows % args.progress_every == 0:
                print(f"scanned_rows={scanned_rows}", file=sys.stderr, flush=True)
            continue
        rn = parse_integer(row[index["rn"]])
        if rn is None or not (
            selection["context_rn_start"] <= rn <= selection["context_rn_end"]
        ):
            continue
        msgdata = row[index["msgdata"]]
        content = row[index["concent"]]
        raw_text, content_source = canonical_content(msgdata, content)
        talker = row[index["strtalker"]]
        if isinstance(msgdata, str):
            stripped = msgdata.strip()
            if stripped.startswith("{") and stripped.endswith("}"):
                try:
                    parsed_msgdata = json.loads(stripped)
                except (json.JSONDecodeError, TypeError, ValueError):
                    parsed_msgdata = None
                if isinstance(parsed_msgdata, dict) and not is_blank(parsed_msgdata.get("strTalker")):
                    talker = parsed_msgdata.get("strTalker")
        raw_rows_by_sample[selection["sample_id"]].append(
            {
                "rn": rn,
                "msgid": row[index["msgid"]],
                "replymsgid": row[index["replymsgid"]],
                "sourceuid": row[index["sourceuid"]],
                "targetuids": row[index["targetuids"]],
                "timestamp": str(row[index["from_unixtime"]]),
                "exported_user_type": str(row[index["user_type"]]),
                "talker": "" if is_blank(talker) else str(talker),
                "raw_text": raw_text,
                "content_source": content_source,
            }
        )
        if args.progress_every and scanned_rows % args.progress_every == 0:
            print(f"scanned_rows={scanned_rows}", file=sys.stderr, flush=True)

    workbook.close()

    output_messages: list[dict[str, Any]] = []
    sample_quality: list[dict[str, Any]] = []
    redaction_totals = Counter()

    for sample_id, selection in sorted(public_by_sample.items()):
        raw_rows = sorted(raw_rows_by_sample.get(sample_id, []), key=lambda item: item["rn"])
        source_order: list[Any] = []
        for item in raw_rows:
            if item["sourceuid"] not in source_order:
                source_order.append(item["sourceuid"])
        actor_by_source = {
            source_id: pseudonym(position + 1)
            for position, source_id in enumerate(source_order)
        }
        replacements: dict[str, str] = {}
        for item in raw_rows:
            actor = actor_by_source[item["sourceuid"]]
            replacements[str(item["sourceuid"])] = actor
            if item["talker"]:
                replacements[item["talker"]] = actor
        msgid_to_rn = {item["msgid"]: item["rn"] for item in raw_rows}

        core_count = 0
        context_count = 0
        for item in raw_rows:
            is_core = selection["core_rn_start"] <= item["rn"] <= selection["core_rn_end"]
            core_count += int(is_core)
            context_count += int(not is_core)
            redacted_text, redaction_counts = redact_text(item["raw_text"], replacements)
            redaction_totals.update(redaction_counts)
            reply_value = item["replymsgid"]
            if reply_value in (None, "", 0, "0"):
                reply_reference = "none"
            elif reply_value in msgid_to_rn:
                reply_reference = f"rn:{msgid_to_rn[reply_value]}"
            else:
                reply_reference = "outside_visible_context"
            exported_role = (
                "未知角色"
                if selection["channel"] == "direct_1v1"
                else item["exported_user_type"]
            )
            output_messages.append(
                {
                    "sample_id": sample_id,
                    "rn": item["rn"],
                    "row_scope": "core" if is_core else "context_only",
                    "timestamp": item["timestamp"],
                    "sender": actor_by_source[item["sourceuid"]],
                    "exported_role": exported_role,
                    "reply_reference": reply_reference,
                    "content_source": item["content_source"],
                    "redacted_text": redacted_text,
                    "text_length": len(redacted_text),
                    "long_text": len(redacted_text) > 1000,
                    "redaction_count": sum(redaction_counts.values()),
                }
            )

        expected_context_count = (
            selection["context_rn_end"] - selection["context_rn_start"] + 1
        )
        sample_quality.append(
            {
                "sample_id": sample_id,
                "expected_rows": expected_context_count,
                "extracted_rows": len(raw_rows),
                "core_rows": core_count,
                "context_only_rows": context_count,
                "status": "ready"
                if len(raw_rows) == expected_context_count and core_count == 20
                else "needs_review",
            }
        )

    result = {
        "classification": "RESTRICTED_RESEARCH_DATA_DO_NOT_COMMIT",
        "dataset_id": "classin-im-pilot0-redacted-review-data-v1-20260829",
        "source_sha256": actual_sha256,
        "public_manifest_id": public_manifest["manifest_id"],
        "selection_method_version": public_manifest.get("selection_method_version"),
        "extraction_method_version": "v1.0",
        "samples": public_manifest["samples"],
        "messages": output_messages,
        "quality": {
            "sample_count": len(sample_quality),
            "message_rows": len(output_messages),
            "core_rows": sum(item["core_rows"] for item in sample_quality),
            "context_only_rows": sum(item["context_only_rows"] for item in sample_quality),
            "ready_samples": sum(item["status"] == "ready" for item in sample_quality),
            "needs_review_samples": sum(
                item["status"] == "needs_review" for item in sample_quality
            ),
            "sample_checks": sample_quality,
            "redaction_totals": dict(redaction_totals),
        },
        "privacy_warning": (
            "Automatic redaction removes common direct identifiers and known participant names, "
            "but free text can retain contextual personal information. Keep this file restricted."
        ),
        "generated_in_seconds": round(time.time() - started, 2),
    }

    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8")
    print(f"output={args.output}")
    print(f"samples={result['quality']['sample_count']}")
    print(f"core_rows={result['quality']['core_rows']}")
    print(f"context_only_rows={result['quality']['context_only_rows']}")
    print(f"needs_review_samples={result['quality']['needs_review_samples']}")


if __name__ == "__main__":
    main()
