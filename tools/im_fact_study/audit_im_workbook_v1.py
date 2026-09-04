#!/usr/bin/env python3
"""Read-only P1 structural and data-quality audit for the ClassIn IM workbook.

The script never emits message text, participant names, user IDs, cluster IDs,
message IDs, or target IDs. It writes only schema information, aggregate counts,
non-reversible structural classifications, and integrity metrics.
"""

from __future__ import annotations

import argparse
import hashlib
import json
import math
import re
import sys
import time
from collections import Counter, defaultdict
from dataclasses import dataclass, field
from datetime import date, datetime, timezone
from pathlib import Path
from typing import Any, Iterable
from zoneinfo import ZoneInfo

from openpyxl import load_workbook


EXPECTED_SHA256 = "a28f4c3125326c2e0f9086f8c0f67104671b94b2992a0173f4c34f90690fee6c"
SHANGHAI = ZoneInfo("Asia/Shanghai")

SAFE_CATEGORICAL_FIELDS = (
    "identity",
    "user_type",
    "user_num",
    "clustertype",
    "msgcmd",
    "dt",
)

SENSITIVE_UNIQUE_FIELDS = (
    "id",
    "clusterid",
    "msgbucketid",
    "msgid",
    "sourceuid",
    "targetuids",
    "replymsgid",
    "strtalker",
)

CONTENT_FIELDS = ("msgdata", "concent")


def field_key(header: Any) -> str:
    text = "" if header is None else str(header).strip()
    return re.split(r"[（(]", text, maxsplit=1)[0].strip().lower()


def is_blank(value: Any) -> bool:
    return value is None or value == ""


def is_empty_reference(value: Any) -> bool:
    if is_blank(value):
        return True
    if isinstance(value, (int, float)) and value == 0:
        return True
    return str(value).strip().lower() in {"", "0", "none", "null", "nan"}


def json_scalar(value: Any) -> Any:
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    if isinstance(value, float) and (math.isnan(value) or math.isinf(value)):
        return str(value)
    return value


def value_type(value: Any) -> str:
    if value is None:
        return "null"
    if isinstance(value, bool):
        return "boolean"
    if isinstance(value, datetime):
        return "datetime"
    if isinstance(value, date):
        return "date"
    if isinstance(value, int):
        return "integer"
    if isinstance(value, float):
        return "float"
    if isinstance(value, str):
        return "string"
    return type(value).__name__


def normalize_category(value: Any) -> str:
    if is_blank(value):
        return "<EMPTY>"
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(json_scalar(value)).strip()


def normalize_channel(value: Any) -> str:
    normalized = normalize_category(value)
    if normalized == "1":
        return "direct_1v1"
    if normalized == "0":
        return "class_group_declared"
    if normalized == "<EMPTY>":
        return "unknown_empty"
    return "unknown_other"


def content_shape(value: Any) -> str:
    if is_blank(value):
        return "EMPTY"
    text = str(value).strip()
    if not text:
        return "WHITESPACE_ONLY"
    lowered = text.lower()
    if lowered.startswith(("http://", "https://")):
        return "URL"
    if text.startswith("{") and text.endswith("}"):
        return "JSON_OBJECT_LIKE"
    if text.startswith("[") and text.endswith("]"):
        return "JSON_ARRAY_LIKE"
    if text.startswith("<") and text.endswith(">"):
        return "MARKUP_LIKE"
    has_cjk = bool(re.search(r"[\u4e00-\u9fff]", text))
    has_latin = bool(re.search(r"[A-Za-z]", text))
    has_digit = bool(re.search(r"\d", text))
    if has_cjk and has_latin:
        return "CJK_LATIN_MIXED"
    if has_cjk:
        return "CJK_TEXT"
    if has_latin:
        return "LATIN_TEXT"
    if has_digit:
        return "NUMERIC_OR_SYMBOL"
    return "SYMBOL_OR_EMOJI"


def target_shape(value: Any) -> str:
    if is_blank(value):
        return "EMPTY"
    text = str(value).strip()
    if not text:
        return "WHITESPACE_ONLY"
    if text.startswith("[") and text.endswith("]"):
        return "ARRAY_LIKE"
    if "," in text:
        return "COMMA_LIST_LIKE"
    return "SCALAR_LIKE"


def length_bucket(length: int) -> str:
    if length == 0:
        return "0"
    if length <= 5:
        return "1-5"
    if length <= 20:
        return "6-20"
    if length <= 50:
        return "21-50"
    if length <= 100:
        return "51-100"
    if length <= 300:
        return "101-300"
    if length <= 1000:
        return "301-1000"
    return "1001+"


def span_bucket(seconds: float | None) -> str:
    if seconds is None:
        return "UNKNOWN"
    if seconds < 0:
        return "NEGATIVE_INVALID"
    if seconds == 0:
        return "0"
    if seconds <= 3600:
        return "<=1h"
    if seconds <= 86400:
        return "1h-1d"
    if seconds <= 7 * 86400:
        return "1d-7d"
    if seconds <= 27 * 86400:
        return "7d-27d"
    return ">27d"


def parse_integer(value: Any) -> int | None:
    if value is None or value == "":
        return None
    if isinstance(value, bool):
        return None
    if isinstance(value, int):
        return value
    if isinstance(value, float) and value.is_integer():
        return int(value)
    text = str(value).strip()
    if re.fullmatch(r"[-+]?\d+", text):
        return int(text)
    return None


def parse_epoch(value: Any) -> float | None:
    if value is None or value == "" or isinstance(value, bool):
        return None
    try:
        parsed = float(value)
    except (TypeError, ValueError):
        return None
    if not math.isfinite(parsed):
        return None
    return parsed


def parse_date(value: Any) -> date | None:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if value is None or value == "":
        return None
    text = str(value).strip()
    for pattern in ("%Y-%m-%d", "%Y/%m/%d", "%Y%m%d"):
        try:
            return datetime.strptime(text, pattern).date()
        except ValueError:
            continue
    return None


def parse_datetime_text(value: Any) -> datetime | None:
    if isinstance(value, datetime):
        return value.replace(tzinfo=value.tzinfo or SHANGHAI)
    if value is None or value == "":
        return None
    text = str(value).strip()
    for pattern in ("%Y-%m-%d %H:%M:%S", "%Y/%m/%d %H:%M:%S"):
        try:
            return datetime.strptime(text, pattern).replace(tzinfo=SHANGHAI)
        except ValueError:
            continue
    return None


def stable_row_digest(values: Iterable[Any]) -> bytes:
    digest = hashlib.blake2b(digest_size=16)
    for value in values:
        if isinstance(value, (datetime, date)):
            token = value.isoformat()
        elif value is None:
            token = "<NULL>"
        else:
            token = str(value)
        digest.update(token.encode("utf-8", errors="replace"))
        digest.update(b"\x1f")
    return digest.digest()


def file_sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def quantiles_from_counter(counter: Counter[int], probabilities: Iterable[float]) -> dict[str, int | None]:
    total = sum(counter.values())
    if not total:
        return {f"p{int(probability * 100):02d}": None for probability in probabilities}
    sorted_items = sorted(counter.items())
    results: dict[str, int | None] = {}
    for probability in probabilities:
        target = max(1, math.ceil(total * probability))
        cumulative = 0
        selected = sorted_items[-1][0]
        for value, count in sorted_items:
            cumulative += count
            if cumulative >= target:
                selected = value
                break
        results[f"p{int(probability * 100):02d}"] = selected
    return results


@dataclass
class ClusterState:
    rows: int = 0
    channels: set[str] = field(default_factory=set)
    user_types: set[str] = field(default_factory=set)
    identities: set[str] = field(default_factory=set)
    user_num_values: set[str] = field(default_factory=set)
    sender_ids: set[Any] = field(default_factory=set)
    rn_values: set[int] = field(default_factory=set)
    rn_invalid: int = 0
    min_rn: int | None = None
    max_rn: int | None = None
    min_epoch: float | None = None
    max_epoch: float | None = None
    rn1_epoch: float | None = None
    rn100_epoch: float | None = None
    date_values: set[str] = field(default_factory=set)

    def add_rn(self, value: int | None) -> None:
        if value is None:
            self.rn_invalid += 1
            return
        self.rn_values.add(value)
        self.min_rn = value if self.min_rn is None else min(self.min_rn, value)
        self.max_rn = value if self.max_rn is None else max(self.max_rn, value)

    def add_epoch(self, value: float | None) -> None:
        if value is None:
            return
        self.min_epoch = value if self.min_epoch is None else min(self.min_epoch, value)
        self.max_epoch = value if self.max_epoch is None else max(self.max_epoch, value)

    def add_ranked_epoch(self, rn_value: int | None, epoch_value: float | None) -> None:
        if rn_value == 1:
            self.rn1_epoch = epoch_value
        elif rn_value == 100:
            self.rn100_epoch = epoch_value


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("input_xlsx", type=Path)
    parser.add_argument("--output", type=Path, required=True)
    parser.add_argument("--expected-sha256", default=EXPECTED_SHA256)
    parser.add_argument("--progress-every", type=int, default=100_000)
    args = parser.parse_args()

    started = time.time()
    actual_sha256 = file_sha256(args.input_xlsx)
    if args.expected_sha256 and actual_sha256 != args.expected_sha256:
        raise SystemExit(
            f"Input fingerprint mismatch: expected {args.expected_sha256}, got {actual_sha256}"
        )

    workbook = load_workbook(args.input_xlsx, read_only=True, data_only=False)
    if not workbook.worksheets:
        raise SystemExit("Workbook contains no worksheets")

    workbook_sheet_summary: list[dict[str, Any]] = []
    for sheet in workbook.worksheets:
        workbook_sheet_summary.append(
            {
                "name": sheet.title,
                "declared_max_row": sheet.max_row,
                "declared_max_column": sheet.max_column,
            }
        )

    worksheet = workbook.worksheets[0]
    rows = worksheet.iter_rows(values_only=True)
    raw_headers = next(rows)
    keys = [field_key(value) for value in raw_headers]
    duplicate_header_keys = sorted(
        key for key, count in Counter(keys).items() if key and count > 1
    )
    index = {key: position for position, key in enumerate(keys)}

    field_stats: dict[str, dict[str, Any]] = {
        key: {
            "blank": 0,
            "whitespace_only": 0,
            "nonempty": 0,
            "type_counts": Counter(),
        }
        for key in keys
    }
    categorical_counts: dict[str, Counter[str]] = {
        key: Counter() for key in SAFE_CATEGORICAL_FIELDS if key in index
    }
    distinct_values: dict[str, set[Any]] = {
        key: set() for key in SENSITIVE_UNIQUE_FIELDS if key in index
    }
    content_shapes: dict[str, Counter[str]] = {
        key: Counter() for key in CONTENT_FIELDS if key in index
    }
    content_length_buckets: dict[str, Counter[str]] = {
        key: Counter() for key in CONTENT_FIELDS if key in index
    }
    content_exact_lengths: dict[str, Counter[int]] = {
        key: Counter() for key in CONTENT_FIELDS if key in index
    }
    target_shapes = Counter()
    target_length_buckets = Counter()
    target_shapes_by_channel: dict[str, Counter[str]] = defaultdict(Counter)

    rows_by_channel = Counter()
    roles_by_channel: dict[str, Counter[str]] = defaultdict(Counter)
    identity_by_channel: dict[str, Counter[str]] = defaultdict(Counter)
    user_num_by_channel: dict[str, Counter[str]] = defaultdict(Counter)
    role_identity_cross = Counter()
    clusters: dict[Any, ClusterState] = {}

    message_ids: set[Any] = set()
    record_ids: set[Any] = set()
    cluster_message_pairs: set[tuple[Any, Any]] = set()
    reply_reference_pairs: list[tuple[Any, Any]] = []
    message_id_duplicates = 0
    record_id_duplicates = 0
    cluster_message_pair_duplicates = 0

    exact_row_digests: set[bytes] = set()
    exact_row_duplicate_occurrences = 0
    payload_digests: set[bytes] = set()
    payload_equivalent_occurrences = 0

    rn_invalid_rows = 0
    epoch_invalid_rows = 0
    dt_invalid_rows = 0
    epoch_dt_comparable_rows = 0
    epoch_dt_mismatch_rows = 0
    timeformat_timetag_comparison = Counter()
    timeformat_timetag_scale_candidates: dict[str, Counter[str]] = {
        "timetag_div_1000": Counter(),
        "timetag_div_1000000": Counter(),
        "timetag_div_1000000000": Counter(),
    }
    timetag_min: float | None = None
    timetag_max: float | None = None
    from_unixtime_comparison = Counter()
    epoch_min: float | None = None
    epoch_max: float | None = None

    content_relationship = Counter()
    reply_reference_state = Counter()
    semantic_placeholder_counts: dict[str, Counter[str]] = defaultdict(Counter)
    msgdata_json_profile = Counter()
    msgdata_json_top_level_keys = Counter()
    msgdata_json_top_level_key_types: dict[str, Counter[str]] = defaultdict(Counter)
    msgdata_concent_equal_top_level_key = Counter()
    parsed_content_relationship = Counter()
    parsed_content_length_delta_buckets = Counter()
    parsed_strtalker_relationship = Counter()

    total_rows = 0
    for row in rows:
        total_rows += 1
        if len(row) < len(keys):
            row = tuple(row) + (None,) * (len(keys) - len(row))

        for position, key in enumerate(keys):
            value = row[position]
            stats = field_stats[key]
            stats["type_counts"][value_type(value)] += 1
            if is_blank(value):
                stats["blank"] += 1
            elif isinstance(value, str) and not value.strip():
                stats["whitespace_only"] += 1
            else:
                stats["nonempty"] += 1
            if key not in CONTENT_FIELDS and isinstance(value, str):
                stripped = value.strip()
                if stripped.lower() in {"null", "none", "nan", "n/a", "unknown"}:
                    semantic_placeholder_counts[key][stripped] += 1

        for key, counter in categorical_counts.items():
            counter[normalize_category(row[index[key]])] += 1

        for key, values in distinct_values.items():
            value = row[index[key]]
            if key == "replymsgid":
                if not is_empty_reference(value):
                    values.add(value)
            elif not is_blank(value):
                values.add(value)

        for key in content_shapes:
            value = row[index[key]]
            content_shapes[key][content_shape(value)] += 1
            length = len(str(value)) if value is not None else 0
            content_length_buckets[key][length_bucket(length)] += 1
            content_exact_lengths[key][length] += 1

        if "targetuids" in index:
            target_value = row[index["targetuids"]]
            target_shapes[target_shape(target_value)] += 1
            target_length = len(str(target_value)) if target_value is not None else 0
            target_length_buckets[length_bucket(target_length)] += 1

        channel = normalize_channel(row[index["clustertype"]]) if "clustertype" in index else "unknown_missing_field"
        rows_by_channel[channel] += 1
        if "targetuids" in index:
            target_shapes_by_channel[channel][target_shape(row[index["targetuids"]])] += 1
        user_type = normalize_category(row[index["user_type"]]) if "user_type" in index else "<MISSING_FIELD>"
        identity = normalize_category(row[index["identity"]]) if "identity" in index else "<MISSING_FIELD>"
        user_num = normalize_category(row[index["user_num"]]) if "user_num" in index else "<MISSING_FIELD>"
        roles_by_channel[channel][user_type] += 1
        identity_by_channel[channel][identity] += 1
        user_num_by_channel[channel][user_num] += 1
        role_identity_cross[(channel, user_type, identity)] += 1

        cluster_id = row[index["clusterid"]] if "clusterid" in index else None
        if not is_blank(cluster_id):
            state = clusters.setdefault(cluster_id, ClusterState())
            state.rows += 1
            state.channels.add(channel)
            state.user_types.add(user_type)
            state.identities.add(identity)
            state.user_num_values.add(user_num)
            if "sourceuid" in index:
                sender = row[index["sourceuid"]]
                if not is_blank(sender):
                    state.sender_ids.add(sender)
            rn_value = parse_integer(row[index["rn"]]) if "rn" in index else None
            state.add_rn(rn_value)
            epoch_value = parse_epoch(row[index["timeformat"]]) if "timeformat" in index else None
            state.add_epoch(epoch_value)
            state.add_ranked_epoch(rn_value, epoch_value)
            if "dt" in index:
                dt_value = parse_date(row[index["dt"]])
                if dt_value is not None:
                    state.date_values.add(dt_value.isoformat())

        rn_value = parse_integer(row[index["rn"]]) if "rn" in index else None
        if "rn" in index and rn_value is None:
            rn_invalid_rows += 1

        epoch_value = parse_epoch(row[index["timeformat"]]) if "timeformat" in index else None
        if "timeformat" in index:
            if epoch_value is None:
                epoch_invalid_rows += 1
            else:
                epoch_min = epoch_value if epoch_min is None else min(epoch_min, epoch_value)
                epoch_max = epoch_value if epoch_max is None else max(epoch_max, epoch_value)

        dt_value = parse_date(row[index["dt"]]) if "dt" in index else None
        if "dt" in index and dt_value is None:
            dt_invalid_rows += 1
        if epoch_value is not None and dt_value is not None:
            epoch_dt_comparable_rows += 1
            epoch_date = datetime.fromtimestamp(epoch_value, tz=timezone.utc).astimezone(SHANGHAI).date()
            if epoch_date != dt_value:
                epoch_dt_mismatch_rows += 1

        timetag_value = parse_epoch(row[index["timetag"]]) if "timetag" in index else None
        if timetag_value is not None:
            timetag_min = timetag_value if timetag_min is None else min(timetag_min, timetag_value)
            timetag_max = timetag_value if timetag_max is None else max(timetag_max, timetag_value)
        if "timetag" in index and "timeformat" in index:
            if timetag_value is None or epoch_value is None:
                timeformat_timetag_comparison["not_comparable"] += 1
            else:
                difference = abs(timetag_value - epoch_value)
                if difference == 0:
                    timeformat_timetag_comparison["exact_equal"] += 1
                elif difference < 1:
                    timeformat_timetag_comparison["within_1_second"] += 1
                elif difference <= 60:
                    timeformat_timetag_comparison["within_60_seconds"] += 1
                else:
                    timeformat_timetag_comparison["over_60_seconds"] += 1
                for candidate_name, divisor in (
                    ("timetag_div_1000", 1_000),
                    ("timetag_div_1000000", 1_000_000),
                    ("timetag_div_1000000000", 1_000_000_000),
                ):
                    candidate_difference = abs((timetag_value / divisor) - epoch_value)
                    if candidate_difference < 1:
                        timeformat_timetag_scale_candidates[candidate_name]["within_1_second"] += 1
                    elif candidate_difference <= 60:
                        timeformat_timetag_scale_candidates[candidate_name]["within_60_seconds"] += 1
                    else:
                        timeformat_timetag_scale_candidates[candidate_name]["over_60_seconds"] += 1

        exported_datetime = (
            parse_datetime_text(row[index["from_unixtime"]]) if "from_unixtime" in index else None
        )
        if "from_unixtime" in index and "timeformat" in index:
            if exported_datetime is None or epoch_value is None:
                from_unixtime_comparison["not_comparable"] += 1
            else:
                epoch_datetime = datetime.fromtimestamp(epoch_value, tz=timezone.utc).astimezone(SHANGHAI)
                if int(exported_datetime.timestamp()) == int(epoch_datetime.timestamp()):
                    from_unixtime_comparison["exact_to_second"] += 1
                else:
                    from_unixtime_comparison["mismatch"] += 1

        record_id = row[index["id"]] if "id" in index else None
        if not is_blank(record_id):
            if record_id in record_ids:
                record_id_duplicates += 1
            else:
                record_ids.add(record_id)

        message_id = row[index["msgid"]] if "msgid" in index else None
        if not is_blank(message_id):
            if message_id in message_ids:
                message_id_duplicates += 1
            else:
                message_ids.add(message_id)
            if not is_blank(cluster_id):
                pair = (cluster_id, message_id)
                if pair in cluster_message_pairs:
                    cluster_message_pair_duplicates += 1
                else:
                    cluster_message_pairs.add(pair)

        if "replymsgid" in index:
            reply_value = row[index["replymsgid"]]
            if is_empty_reference(reply_value):
                reply_reference_state["empty_or_zero"] += 1
            else:
                reply_reference_state["nonempty"] += 1
                if not is_blank(cluster_id):
                    reply_reference_pairs.append((cluster_id, reply_value))
                else:
                    reply_reference_state["nonempty_without_cluster"] += 1

        exact_digest = stable_row_digest(row[: len(keys)])
        if exact_digest in exact_row_digests:
            exact_row_duplicate_occurrences += 1
        else:
            exact_row_digests.add(exact_digest)

        payload_fields = (
            "clusterid",
            "sourceuid",
            "targetuids",
            "timeformat",
            "msgcmd",
            "msgdata",
            "concent",
        )
        payload_digest = stable_row_digest(
            row[index[key]] if key in index else None for key in payload_fields
        )
        if payload_digest in payload_digests:
            payload_equivalent_occurrences += 1
        else:
            payload_digests.add(payload_digest)

        if "msgdata" in index and "concent" in index:
            msgdata = row[index["msgdata"]]
            content = row[index["concent"]]
            msgdata_blank = is_blank(msgdata) or (isinstance(msgdata, str) and not msgdata.strip())
            content_blank = is_blank(content) or (isinstance(content, str) and not content.strip())
            if msgdata_blank and content_blank:
                content_relationship["both_empty"] += 1
            elif msgdata_blank:
                content_relationship["only_concent_nonempty"] += 1
            elif content_blank:
                content_relationship["only_msgdata_nonempty"] += 1
            elif str(msgdata) == str(content):
                content_relationship["exact_equal"] += 1
            else:
                content_relationship["both_nonempty_different"] += 1

            if isinstance(msgdata, str):
                stripped_msgdata = msgdata.strip()
                if stripped_msgdata.startswith("{") and stripped_msgdata.endswith("}"):
                    msgdata_json_profile["object_like"] += 1
                    try:
                        parsed_msgdata = json.loads(stripped_msgdata)
                    except (json.JSONDecodeError, TypeError, ValueError):
                        msgdata_json_profile["object_like_parse_failed"] += 1
                    else:
                        if isinstance(parsed_msgdata, dict):
                            msgdata_json_profile["valid_top_level_object"] += 1
                            for json_key, json_value in parsed_msgdata.items():
                                safe_key = str(json_key)[:120]
                                msgdata_json_top_level_keys[safe_key] += 1
                                msgdata_json_top_level_key_types[safe_key][value_type(json_value)] += 1
                                if not content_blank and str(json_value) == str(content):
                                    msgdata_concent_equal_top_level_key[safe_key] += 1

                            parsed_content = parsed_msgdata.get("content")
                            parsed_content_blank = is_blank(parsed_content) or (
                                isinstance(parsed_content, str) and not parsed_content.strip()
                            )
                            if parsed_content_blank and content_blank:
                                parsed_content_relationship["both_empty"] += 1
                            elif parsed_content_blank:
                                parsed_content_relationship["only_concent_nonempty"] += 1
                            elif content_blank:
                                parsed_content_relationship["only_parsed_content_nonempty"] += 1
                            elif str(parsed_content) == str(content):
                                parsed_content_relationship["exact_equal"] += 1
                            elif str(parsed_content).strip() == str(content).strip():
                                parsed_content_relationship["equal_after_strip"] += 1
                            else:
                                parsed_content_relationship["different"] += 1
                                delta = abs(len(str(parsed_content)) - len(str(content)))
                                parsed_content_length_delta_buckets[length_bucket(delta)] += 1

                            parsed_talker = parsed_msgdata.get("strTalker")
                            exported_talker = row[index["strtalker"]] if "strtalker" in index else None
                            parsed_talker_blank = is_blank(parsed_talker) or (
                                isinstance(parsed_talker, str) and not parsed_talker.strip()
                            )
                            exported_talker_blank = is_blank(exported_talker) or (
                                isinstance(exported_talker, str) and not exported_talker.strip()
                            )
                            if parsed_talker_blank and exported_talker_blank:
                                parsed_strtalker_relationship["both_empty"] += 1
                            elif parsed_talker_blank:
                                parsed_strtalker_relationship["only_exported_nonempty"] += 1
                            elif exported_talker_blank:
                                parsed_strtalker_relationship["only_parsed_nonempty"] += 1
                            elif str(parsed_talker) == str(exported_talker):
                                parsed_strtalker_relationship["exact_equal"] += 1
                            else:
                                parsed_strtalker_relationship["different"] += 1
                        else:
                            msgdata_json_profile["valid_non_object"] += 1
                else:
                    msgdata_json_profile["not_object_like"] += 1

        if args.progress_every and total_rows % args.progress_every == 0:
            print(f"scanned_rows={total_rows}", file=sys.stderr, flush=True)

    reply_resolved_same_cluster = sum(
        1 for pair in reply_reference_pairs if pair in cluster_message_pairs
    )
    reply_reference_state["resolved_in_same_cluster_export_window"] = reply_resolved_same_cluster
    reply_reference_state["unresolved_in_same_cluster_export_window"] = (
        len(reply_reference_pairs) - reply_resolved_same_cluster
    )

    cluster_size_distribution = Counter()
    cluster_channel_distribution = Counter()
    cluster_role_composition: dict[str, Counter[str]] = defaultdict(Counter)
    cluster_identity_composition: dict[str, Counter[str]] = defaultdict(Counter)
    cluster_user_num_consistency = Counter()
    cluster_channel_consistency = Counter()
    cluster_rn_quality = Counter()
    cluster_sender_count_distribution: dict[str, Counter[int]] = defaultdict(Counter)
    cluster_date_count_distribution: dict[str, Counter[int]] = defaultdict(Counter)
    cluster_span_distribution: dict[str, Counter[str]] = defaultdict(Counter)
    cluster_user_num_distribution: dict[str, Counter[str]] = defaultdict(Counter)
    cluster_rank_time_direction: dict[str, Counter[str]] = defaultdict(Counter)
    observed_group_role_summary = Counter()

    for state in clusters.values():
        channel_label = next(iter(state.channels)) if len(state.channels) == 1 else "mixed_or_unknown"
        cluster_size_distribution[state.rows] += 1
        cluster_channel_distribution[channel_label] += 1
        cluster_channel_consistency["single_channel"] += int(len(state.channels) == 1)
        cluster_channel_consistency["mixed_channel"] += int(len(state.channels) > 1)
        role_key = " | ".join(sorted(state.user_types))
        identity_key = " | ".join(sorted(state.identities))
        cluster_role_composition[channel_label][role_key] += 1
        cluster_identity_composition[channel_label][identity_key] += 1
        for user_num_value in state.user_num_values:
            cluster_user_num_distribution[channel_label][user_num_value] += 1
        cluster_user_num_consistency["single_value"] += int(len(state.user_num_values) == 1)
        cluster_user_num_consistency["multiple_values"] += int(len(state.user_num_values) > 1)
        cluster_sender_count_distribution[channel_label][len(state.sender_ids)] += 1
        cluster_date_count_distribution[channel_label][len(state.date_values)] += 1

        if state.rn_invalid:
            cluster_rn_quality["contains_invalid_or_missing_rn"] += 1
        elif state.rn_values == set(range(1, 101)) and state.rows == 100:
            cluster_rn_quality["exact_1_to_100_and_100_rows"] += 1
        elif state.min_rn == 1 and state.max_rn == state.rows and len(state.rn_values) == state.rows:
            cluster_rn_quality["contiguous_1_to_row_count"] += 1
        elif len(state.rn_values) < state.rows:
            cluster_rn_quality["duplicate_rn_within_cluster"] += 1
        else:
            cluster_rn_quality["other_noncontiguous"] += 1

        span = None
        if state.min_epoch is not None and state.max_epoch is not None:
            span = state.max_epoch - state.min_epoch
        cluster_span_distribution[channel_label][span_bucket(span)] += 1

        if state.rn1_epoch is None or state.rn100_epoch is None:
            cluster_rank_time_direction[channel_label]["endpoint_time_missing"] += 1
        elif state.rn1_epoch > state.rn100_epoch:
            cluster_rank_time_direction[channel_label]["rn1_newer_than_rn100"] += 1
        elif state.rn1_epoch < state.rn100_epoch:
            cluster_rank_time_direction[channel_label]["rn1_older_than_rn100"] += 1
        else:
            cluster_rank_time_direction[channel_label]["same_endpoint_second"] += 1

        if channel_label == "class_group_declared":
            learner_roles = {"学生", "旁听生"}
            staff_roles = {"班主任", "教师"}
            has_learner = bool(state.user_types & learner_roles)
            has_staff = bool(state.user_types & staff_roles)
            if has_learner and has_staff:
                observed_group_role_summary["learner_and_staff_observed"] += 1
            elif has_learner:
                observed_group_role_summary["learner_only_observed"] += 1
            elif has_staff and state.user_types <= staff_roles:
                observed_group_role_summary["staff_only_observed"] += 1
            else:
                observed_group_role_summary["other_or_unknown_observed"] += 1

    epoch_min_iso = None
    epoch_max_iso = None
    if epoch_min is not None:
        epoch_min_iso = datetime.fromtimestamp(epoch_min, tz=timezone.utc).astimezone(SHANGHAI).isoformat()
    if epoch_max is not None:
        epoch_max_iso = datetime.fromtimestamp(epoch_max, tz=timezone.utc).astimezone(SHANGHAI).isoformat()

    field_summary = []
    for position, key in enumerate(keys):
        stats = field_stats[key]
        field_summary.append(
            {
                "position": position + 1,
                "key": key,
                "raw_header": str(raw_headers[position]),
                "blank": stats["blank"],
                "whitespace_only": stats["whitespace_only"],
                "nonempty": stats["nonempty"],
                "missing_rate": round((stats["blank"] + stats["whitespace_only"]) / total_rows, 6)
                if total_rows
                else None,
                "type_counts": dict(stats["type_counts"].most_common()),
                "sensitivity": "restricted_content"
                if key in CONTENT_FIELDS
                else "restricted_identifier"
                if key in SENSITIVE_UNIQUE_FIELDS
                else "internal_metadata",
                "distinct_nonempty": len(distinct_values[key]) if key in distinct_values else None,
            }
        )

    result = {
        "audit": {
            "audit_id": "classin-im-p1-structure-audit-v1-2-20260829",
            "method_version": "v1.2",
            "generated_at": datetime.now(tz=SHANGHAI).isoformat(),
            "elapsed_seconds": round(time.time() - started, 2),
            "privacy": "aggregate-only; no message text or identifiers emitted",
        },
        "source": {
            "file_name": args.input_xlsx.name,
            "file_size_bytes": args.input_xlsx.stat().st_size,
            "sha256": actual_sha256,
            "sheet_count": len(workbook.worksheets),
            "sheets": workbook_sheet_summary,
            "active_sheet_audited": worksheet.title,
            "data_rows_scanned": total_rows,
        },
        "schema": {
            "duplicate_normalized_header_keys": duplicate_header_keys,
            "fields": field_summary,
        },
        "categorical_counts": {
            key: dict(counter.most_common()) for key, counter in categorical_counts.items()
        },
        "channel_and_role": {
            "rows_by_declared_channel": dict(rows_by_channel.most_common()),
            "user_type_by_channel": {
                channel: dict(counter.most_common()) for channel, counter in roles_by_channel.items()
            },
            "identity_by_channel": {
                channel: dict(counter.most_common()) for channel, counter in identity_by_channel.items()
            },
            "user_num_by_channel": {
                channel: dict(counter.most_common()) for channel, counter in user_num_by_channel.items()
            },
            "user_type_identity_cross": [
                {
                    "channel": channel,
                    "user_type": user_type,
                    "identity": identity,
                    "rows": count,
                }
                for (channel, user_type, identity), count in role_identity_cross.most_common()
            ],
            "observed_cluster_role_composition": {
                channel: dict(counter.most_common()) for channel, counter in cluster_role_composition.items()
            },
            "observed_cluster_identity_composition": {
                channel: dict(counter.most_common()) for channel, counter in cluster_identity_composition.items()
            },
            "observed_group_role_summary": dict(observed_group_role_summary.most_common()),
            "warning": "Observed sender-role composition in the exported rows does not prove full group membership or group purpose.",
        },
        "sampling_and_cluster_structure": {
            "cluster_count": len(clusters),
            "cluster_size_distribution": {
                str(size): count for size, count in sorted(cluster_size_distribution.items())
            },
            "clusters_by_declared_channel": dict(cluster_channel_distribution.most_common()),
            "cluster_channel_consistency": dict(cluster_channel_consistency),
            "cluster_user_num_consistency": dict(cluster_user_num_consistency),
            "cluster_rn_quality": dict(cluster_rn_quality.most_common()),
            "sender_count_distribution_by_channel": {
                channel: {str(size): count for size, count in sorted(counter.items())}
                for channel, counter in cluster_sender_count_distribution.items()
            },
            "date_count_distribution_by_channel": {
                channel: {str(size): count for size, count in sorted(counter.items())}
                for channel, counter in cluster_date_count_distribution.items()
            },
            "time_span_distribution_by_channel": {
                channel: dict(counter.most_common()) for channel, counter in cluster_span_distribution.items()
            },
            "user_num_distribution_by_cluster_and_channel": {
                channel: dict(counter.most_common())
                for channel, counter in cluster_user_num_distribution.items()
            },
            "rank_time_direction_by_channel": {
                channel: dict(counter.most_common())
                for channel, counter in cluster_rank_time_direction.items()
            },
        },
        "identifier_integrity": {
            "record_id_nonempty_distinct": len(record_ids),
            "record_id_duplicate_occurrences": record_id_duplicates,
            "message_id_nonempty_distinct": len(message_ids),
            "message_id_duplicate_occurrences": message_id_duplicates,
            "cluster_message_pair_nonempty_distinct": len(cluster_message_pairs),
            "cluster_message_pair_duplicate_occurrences": cluster_message_pair_duplicates,
            "exact_row_duplicate_occurrences": exact_row_duplicate_occurrences,
            "payload_equivalent_occurrences": payload_equivalent_occurrences,
            "reply_reference_state": dict(reply_reference_state),
            "sensitive_field_distinct_counts": {
                key: len(values) for key, values in distinct_values.items()
            },
            "note": "Distinct counts are emitted without values; payload equivalence is a structural flag, not an instruction to deduplicate business messages.",
        },
        "time_integrity": {
            "rn_invalid_or_missing_rows": rn_invalid_rows,
            "epoch_invalid_or_missing_rows": epoch_invalid_rows,
            "dt_invalid_or_missing_rows": dt_invalid_rows,
            "epoch_dt_comparable_rows": epoch_dt_comparable_rows,
            "epoch_dt_mismatch_rows": epoch_dt_mismatch_rows,
            "timeformat_timetag_comparison": dict(timeformat_timetag_comparison.most_common()),
            "timeformat_timetag_scale_candidates": {
                candidate: dict(counter.most_common())
                for candidate, counter in timeformat_timetag_scale_candidates.items()
            },
            "timetag_min": timetag_min,
            "timetag_max": timetag_max,
            "from_unixtime_comparison": dict(from_unixtime_comparison.most_common()),
            "epoch_min_asia_shanghai": epoch_min_iso,
            "epoch_max_asia_shanghai": epoch_max_iso,
        },
        "content_structure": {
            "shape_counts": {
                key: dict(counter.most_common()) for key, counter in content_shapes.items()
            },
            "length_buckets": {
                key: dict(counter.most_common()) for key, counter in content_length_buckets.items()
            },
            "length_quantiles": {
                key: quantiles_from_counter(counter, (0.50, 0.90, 0.95, 0.99))
                for key, counter in content_exact_lengths.items()
            },
            "max_length": {
                key: max(counter) if counter else None for key, counter in content_exact_lengths.items()
            },
            "msgdata_concent_relationship": dict(content_relationship.most_common()),
            "msgdata_json_profile": dict(msgdata_json_profile.most_common()),
            "msgdata_json_top_level_keys": dict(msgdata_json_top_level_keys.most_common()),
            "msgdata_json_top_level_key_types": {
                key: dict(counter.most_common())
                for key, counter in msgdata_json_top_level_key_types.items()
            },
            "msgdata_concent_equal_top_level_key": dict(
                msgdata_concent_equal_top_level_key.most_common()
            ),
            "parsed_content_vs_concent": dict(parsed_content_relationship.most_common()),
            "parsed_content_length_delta_buckets_when_different": dict(
                parsed_content_length_delta_buckets.most_common()
            ),
            "parsed_strtalker_vs_strtalker": dict(parsed_strtalker_relationship.most_common()),
            "targetuids_shape_counts": dict(target_shapes.most_common()),
            "targetuids_shape_counts_by_channel": {
                channel: dict(counter.most_common())
                for channel, counter in target_shapes_by_channel.items()
            },
            "targetuids_length_buckets": dict(target_length_buckets.most_common()),
        },
        "semantic_placeholders": {
            key: dict(counter.most_common())
            for key, counter in semantic_placeholder_counts.items()
            if counter
        },
        "interpretation_limits": [
            "The workbook contains exported rows, not necessarily complete natural conversations.",
            "A declared class-group flag does not prove the observed messages are teacher-student instruction rather than teaching operations or staff coordination.",
            "Observed sender roles do not prove full membership composition.",
            "A nonempty reply reference does not prove that the issue was understood or resolved.",
            "Payload-equivalent rows may be legitimate repeated communication and must not be silently removed.",
            "No semantic topic, task, need, learning outcome, or AI opportunity is inferred in P1.",
        ],
    }

    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8")
    workbook.close()
    print(f"audit_output={args.output}")
    print(f"elapsed_seconds={result['audit']['elapsed_seconds']}")


if __name__ == "__main__":
    main()
