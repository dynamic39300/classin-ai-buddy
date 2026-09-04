#!/usr/bin/env python3
"""Validate Topic Pilot 50 v1.2 calibration and evidence preservation."""

from __future__ import annotations

import hashlib
import json
import re
import sys
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook


EVIDENCE_SHEETS = ["01_案例目录", "02_案例消息", "06_原表字段索引"]
FORBIDDEN_PHRASES = ["无教学诉求", "无明确教学诉求", "无具体教学内容讨论", "无进一步求助", "未提出求助"]


def sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def sheet_values(ws):
    return tuple(tuple(cell.value for cell in row) for row in ws.iter_rows())


def topic_lines(value):
    return [line.strip() for line in str(value or "").splitlines() if line.strip()]


def main() -> None:
    if len(sys.argv) != 5:
        raise SystemExit("Usage: validate_topic_pilot50_v1_2.py <completed-v1.xlsx> <v1.2.xlsx> <manifest.json> <qa.json>")
    source_path, output_path, manifest_path, qa_path = map(lambda value: Path(value).resolve(), sys.argv[1:])
    manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
    source_wb = load_workbook(source_path, data_only=False)
    output_wb = load_workbook(output_path, data_only=False)
    source_topic = source_wb["03_Topic人工标注"]
    output_topic = output_wb["03_Topic人工标注"]
    source_topics = [source_topic.cell(row, 10).value for row in range(5, 55)]
    output_topics = [output_topic.cell(row, 10).value for row in range(5, 55)]
    row_by_case = {output_topic.cell(row, 2).value: row for row in range(5, 55)}

    expected_counts = {"TP50-B3-DIR-04": 2, "TP50-B3-DIR-05": 4, "TP50-B3-DIR-06": 4, "TP50-B1-GRP-02": 3}
    checks = {
        "sheet_names_preserved": source_wb.sheetnames == output_wb.sheetnames,
        "active_sheet": output_wb.active.title,
        "evidence_sheets_identical": {name: sheet_values(source_wb[name]) == sheet_values(output_wb[name]) for name in EVIDENCE_SHEETS},
        "topic_context_A_to_I_identical": all(
            source_topic.cell(row, col).value == output_topic.cell(row, col).value
            for row in range(4, 55) for col in range(1, 10)
        ),
        "filled_topic_cells": sum(bool(value) for value in output_topics),
        "source_topic_lines": sum(len(topic_lines(value)) for value in source_topics),
        "output_topic_lines": sum(len(topic_lines(value)) for value in output_topics),
        "corrected_case_topic_counts": {
            case_id: len(topic_lines(output_topic.cell(row_by_case[case_id], 10).value)) for case_id in expected_counts
        },
        "forbidden_phrase_counts": {
            phrase: sum(str(value or "").count(phrase) for value in output_topics) for phrase in FORBIDDEN_PHRASES
        },
        "topic_numbering_valid": all(
            [int(match.group(1)) if (match := re.match(r"^T(\d+)\.", line)) else None for line in topic_lines(value)]
            == list(range(1, len(topic_lines(value)) + 1))
            for value in output_topics
        ),
        "topic_sheet_max_column": output_topic.max_column,
        "auto_filter": output_topic.auto_filter.ref,
        "data_validation_count": len(output_topic.data_validations.dataValidation),
        "calibration_records": sum(bool(output_wb["04_问题与校准"].cell(row, 1).value) for row in range(5, 10)),
        "source_sha_matches_manifest": sha256(source_path) == manifest["source_workbook"]["sha256"],
        "output_sha_matches_manifest": sha256(output_path) == manifest["output_workbook"]["sha256"],
    }
    expected = {
        "sheet_names_preserved": True,
        "active_sheet": "03_Topic人工标注",
        "evidence_sheets_identical": {name: True for name in EVIDENCE_SHEETS},
        "topic_context_A_to_I_identical": True,
        "filled_topic_cells": 50,
        "source_topic_lines": 183,
        "output_topic_lines": 183,
        "corrected_case_topic_counts": expected_counts,
        "forbidden_phrase_counts": {phrase: 0 for phrase in FORBIDDEN_PHRASES},
        "topic_numbering_valid": True,
        "topic_sheet_max_column": 10,
        "auto_filter": "A4:J54",
        "data_validation_count": 0,
        "calibration_records": 5,
        "source_sha_matches_manifest": True,
        "output_sha_matches_manifest": True,
    }
    failures = {key: {"actual": checks.get(key), "expected": value} for key, value in expected.items() if checks.get(key) != value}
    report = {
        "status": "PASS" if not failures else "FAIL",
        "validated_at": datetime.now().isoformat(timespec="seconds"),
        "source_workbook": {"path": str(source_path), "sha256": sha256(source_path)},
        "output_workbook": {"path": str(output_path), "sha256": sha256(output_path)},
        "manifest": {"path": str(manifest_path), "sha256": sha256(manifest_path)},
        "checks": checks,
        "failures": failures,
    }
    source_wb.close()
    output_wb.close()
    qa_path.write_text(json.dumps(report, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps(report, ensure_ascii=False, indent=2))
    if failures:
        raise SystemExit(1)


if __name__ == "__main__":
    main()
