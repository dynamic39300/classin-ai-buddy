#!/usr/bin/env python3
"""Validate Topic Pilot 50 v1.1 against the user-edited v1 evidence baseline."""

from __future__ import annotations

import hashlib
import json
import sys
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook


EVIDENCE_SHEETS = ["01_案例目录", "02_案例消息", "06_原表字段索引"]


def sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def sheet_values(ws):
    return tuple(tuple(cell.value for cell in row) for row in ws.iter_rows())


def main() -> None:
    if len(sys.argv) != 5:
        raise SystemExit("Usage: validate_topic_pilot50_v1_1.py <v1.xlsx> <v1.1.xlsx> <manifest.json> <qa.json>")
    source_path, output_path, manifest_path, qa_path = map(lambda item: Path(item).resolve(), sys.argv[1:])
    manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
    source_wb = load_workbook(source_path, data_only=False)
    output_wb = load_workbook(output_path, data_only=False)

    source_topic = source_wb["03_Topic人工标注"]
    output_topic = output_wb["03_Topic人工标注"]
    source_topics = [source_topic.cell(row, 10).value for row in range(5, 55)]
    output_topics = [output_topic.cell(row, 10).value for row in range(5, 55)]
    checks = {
        "sheet_names_preserved": source_wb.sheetnames == output_wb.sheetnames,
        "active_sheet": output_wb.active.title,
        "evidence_sheets_identical": {
            name: sheet_values(source_wb[name]) == sheet_values(output_wb[name]) for name in EVIDENCE_SHEETS
        },
        "topic_context_A_to_I_identical": all(
            source_topic.cell(row, col).value == output_topic.cell(row, col).value
            for row in range(4, 55) for col in range(1, 10)
        ),
        "topic_values_identical": source_topics == output_topics,
        "batch1_filled_topics": sum(bool(value) for value in output_topics[:10]),
        "batch2_3_blank_topics": sum(value in (None, "") for value in output_topics[10:]),
        "topic_sheet_max_column": output_topic.max_column,
        "sole_input_header": output_topic["J4"].value,
        "auto_filter": output_topic.auto_filter.ref,
        "data_validation_count": len(output_topic.data_validations.dataValidation),
        "guide_has_five_utterance_rule": "至少出现5句有效发言" in str(output_wb["00_使用说明"]["B7"].value),
        "audit_has_five_utterance_rule": "≥5句有效发言" in str(output_wb["05_抽样审计"]["B15"].value),
        "source_sha_matches_manifest": sha256(source_path) == manifest["source_workbook"]["sha256"],
        "output_sha_matches_manifest": sha256(output_path) == manifest["output_workbook"]["sha256"],
    }
    expected = {
        "sheet_names_preserved": True,
        "active_sheet": "03_Topic人工标注",
        "evidence_sheets_identical": {name: True for name in EVIDENCE_SHEETS},
        "topic_context_A_to_I_identical": True,
        "topic_values_identical": True,
        "batch1_filled_topics": 10,
        "batch2_3_blank_topics": 40,
        "topic_sheet_max_column": 10,
        "auto_filter": "A4:J54",
        "data_validation_count": 0,
        "guide_has_five_utterance_rule": True,
        "audit_has_five_utterance_rule": True,
        "source_sha_matches_manifest": True,
        "output_sha_matches_manifest": True,
    }
    failures = {
        key: {"actual": checks.get(key), "expected": value}
        for key, value in expected.items() if checks.get(key) != value
    }
    report = {
        "status": "PASS" if not failures else "FAIL",
        "validated_at": datetime.now().isoformat(timespec="seconds"),
        "source_workbook": {"path": str(source_path), "sha256": sha256(source_path)},
        "output_workbook": {"path": str(output_path), "sha256": sha256(output_path)},
        "manifest": {"path": str(manifest_path), "sha256": sha256(manifest_path)},
        "checks": checks,
        "failures": failures,
    }
    source_wb.close()
    output_wb.close()
    qa_path.write_text(json.dumps(report, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps(report, ensure_ascii=False, indent=2))
    if failures:
        raise SystemExit(1)


if __name__ == "__main__":
    main()
