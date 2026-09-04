#!/usr/bin/env python3
"""Create Pilot 0 v2.3.3 workbooks by removing the event-gap field.

The source workbooks are v2.3.2. This script makes no semantic changes to
message evidence or existing annotation values. It physically removes the
event-level gap column, shifts the sample layer left, rewrites the read-only
summary formulas, and removes the same concept from examples and field help.
"""

from __future__ import annotations

import hashlib
import json
import sys
from copy import copy
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.styles.cell_style import StyleArray
from openpyxl.worksheet.datavalidation import DataValidation


def sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def copy_cell(source: Any, target: Any) -> None:
    target.value = source.value
    if source.has_style:
        target._style = copy(source._style)
    if source.number_format:
        target.number_format = source.number_format
    target.font = copy(source.font)
    target.fill = copy(source.fill)
    target.border = copy(source.border)
    target.alignment = copy(source.alignment)
    target.protection = copy(source.protection)


def clear_cell(cell: Any) -> None:
    cell.value = None
    cell._style = StyleArray()
    cell.number_format = "General"


def standard_summary_rows(ws: Any) -> list[int]:
    seen: set[str] = set()
    rows: list[int] = []
    for row in range(8, 678):
        sample_id = str(ws.cell(row, 1).value or "")
        scope = str(ws.cell(row, 7).value or "")
        if sample_id and scope == "核心-需标注" and sample_id not in seen:
            seen.add(sample_id)
            rows.append(row)
    return rows


def add_list_validation(ws: Any, coordinate: str, values: list[str]) -> None:
    joined = ",".join(values)
    dv = DataValidation(type="list", formula1=f'"{joined}"', allow_blank=True)
    dv.promptTitle = "按可见证据选择"
    dv.prompt = " / ".join(values)
    dv.errorTitle = "请使用规定取值"
    dv.error = " / ".join(values)
    dv.errorStyle = "stop"
    dv.showErrorMessage = True
    dv.showInputMessage = True
    dv.showDropDown = False
    ws.add_data_validation(dv)
    dv.add(ws[coordinate])


def remove_event_gap_from_coding(ws: Any) -> list[int]:
    # Remove merges that cross the deleted AG column, then recreate the three
    # section labels after the sample layer shifts from AH:AR to AG:AQ.
    for merged in ["AD6:AG6", "AH6:AR6", "A5:AR5"]:
        if merged in ws.merged_cells:
            ws.unmerge_cells(merged)

    # Save existing validations that are not part of the sample layer.
    retained = []
    for dv in ws.data_validations.dataValidation:
        refs = str(dv.sqref)
        if not any(refs.startswith(prefix) for prefix in ("AH", "AI", "AP", "AQ")):
            retained.append(dv)
    ws.data_validations.dataValidation = retained

    ws.delete_cols(33, 1)  # AG: 事项关键缺口/备注
    ws.merge_cells("A5:AQ5")
    ws.merge_cells("AD6:AF6")
    ws.merge_cells("AG6:AQ6")

    ws["A1"] = ws["A1"].value.replace("v2.3.2", "v2.3.3")
    ws["A2"] = (
        "只填本页：黄色=每条核心消息；绿色=事项首次出现行；紫色=每个 sample_id 唯一样本总结行。"
        "事项层只保留中性名称、对方期待和可见状态；05 自动镜像，不填写。"
    )
    ws["A5"] = (
        "方法审阅版｜灰色只读；黄色消息层；绿色事项层；紫色样本层。"
        "先读完同一 sample_id，再按层级填写；新请求尚无结果时，事项状态标‘未看到’。"
    )
    ws["AD6"] = "事项层｜同一事项首次出现行填写一次；仅记录名称、对方期待与可见状态"
    ws["AG6"] = "样本层｜仅每个 sample_id 的唯一紫色核心行填写一次；05 自动镜像"

    rows = standard_summary_rows(ws)
    for row in rows:
        add_list_validation(ws, f"AG{row}", ["充分", "部分", "不足", "无法判断"])
        add_list_validation(ws, f"AH{row}", ["是", "否", "无法判断"])
        add_list_validation(ws, f"AO{row}", ["高", "中", "低", "无法判断"])
        add_list_validation(
            ws,
            f"AP{row}",
            ["无", "误看第一次分析材料", "看到另一标注者结果", "外部反查身份", "其他"],
        )
    return rows


def update_examples(ws: Any, annotator: str) -> None:
    ws["A1"] = f"开放编码规范示例 v2.3.3 · 标注员 {annotator} · 全部为虚构消息"
    ws["A2"] = (
        "每个填写标题均带取值或推荐写法，并列 Good case / Bad case。"
        "事项层只记录中性名称、对方期待和可见状态；尚未出现结果标为‘未看到’。"
    )

    # Replace the legacy wide event-summary example with the three fields that
    # actually remain in sheet 03. This prevents a removed gap concept from
    # reappearing under a different name.
    compact_headers = [
        "事件ID【E1/E2…】\nGood:同一事项同ID\nBad:预设5个",
        "事项中性名称【事项首行·自由】\nGood:确认开始时间\nBad:需要AI提醒",
        "对方被期待提供什么【事项首行·自由】\nGood:提供开始时间\nBad:改善体验",
        "事项可见状态【事项首行·下拉】\nGood:按完整可见链\nBad:有人回复就完成",
    ]
    compact_rows = [
        ["E1", "确认课程开始时间", "提供明确开始时间", "完成"],
        ["E2", "收集第3页图片", "提交第3页图片", "仅有进展"],
    ]
    for col, value in enumerate(compact_headers, 1):
        ws.cell(17, col).value = value
    for row_index, values in zip((18, 19), compact_rows):
        for col, value in enumerate(values, 1):
            ws.cell(row_index, col).value = value
    for row in range(17, 20):
        for col in range(5, 11):
            clear_cell(ws.cell(row, col))

    # Q was the removed field. Shift the two remaining guidance columns left.
    for row in range(48, 53):
        copy_cell(ws.cell(row, 18), ws.cell(row, 17))
        copy_cell(ws.cell(row, 19), ws.cell(row, 18))
        clear_cell(ws.cell(row, 19))
    ws.column_dimensions["Q"].width = ws.column_dimensions["R"].width
    ws.column_dimensions["R"].width = ws.column_dimensions["S"].width
    ws.column_dimensions["S"].width = 2


def update_field_guide(ws: Any) -> None:
    ws["A1"] = "v2.3.3 字段选项、正反例与一体化填写速查"
    ws["A2"] = (
        "事项层只记录中性名称、对方期待和可见状态。请求刚发起而尚无结果是‘未看到’，"
        "不是缺口；证据限制统一写在消息层‘证据边界/推断警报’。"
    )
    ws.delete_rows(45, 1)


def update_summary_formulas(ws: Any, summary_rows: list[int], coding_ws: Any) -> None:
    ws["A1"] = ws["A1"].value.replace("v2.3.2", "v2.3.3")
    ws["A2"] = (
        "本页自动汇总 03 中每个 sample_id 的唯一紫色样本总结行。"
        "请只在 03 填写；这里不再人工关联 ID、不再录入。"
    )
    row_by_sample = {
        str(coding_ws.cell(row, 1).value): row
        for row in summary_rows
        if coding_ws.cell(row, 1).value
    }
    source_columns = ["AG", "AH", "AI", "AJ", "AK", "AL", "AM", "AN", "AO", "AP", "AQ"]
    for row in range(8, 32):
        sample_id = str(ws.cell(row, 1).value or "")
        source_row = row_by_sample.get(sample_id)
        if source_row is None:
            continue
        for offset, source_column in enumerate(source_columns, start=3):
            ref = f"'03_逐条开放编码'!{source_column}{source_row}"
            ws.cell(row, offset).value = f'=IF({ref}="","",{ref})'


def update_version_text(wb: Any, annotator: str) -> None:
    instructions = wb["00_使用与溯源说明"]
    instructions["A1"] = f"ClassIn IM Pilot 0 字段精简审阅 v2.3.3 · 标注员 {annotator}"
    instructions["A2"] = (
        "v2.3.3 删除事项缺口字段：未出现结果由事项状态‘未看到’表达；"
        "数据或证据限制统一记录在消息层‘证据边界/推断警报’。"
    )
    instructions["A34"] = "先看 00A 示例表头；事项层现在只保留中性名称、对方期待和可见状态。"

    mapping = wb["01_样本与会话映射"]
    mapping["A1"] = "Pilot 0 样本、会话与角色上下文映射 v2.3.3"
    mapping["A2"] = (
        "标准轨所有人工填写集中在 03；事项层已删除缺口字段。"
        "05 自动镜像 03 的样本层；M:O 是样本可见发言人与角色的机械汇总。"
    )

    wb["08_单方可见汇总"]["A1"] = f"窗口内单方可见样本汇总 v2.3.3 · 标注员 {annotator}"


def transform(input_path: Path, output_path: Path, annotator: str) -> dict[str, Any]:
    wb = load_workbook(input_path)
    coding = wb["03_逐条开放编码"]
    summary_rows = remove_event_gap_from_coding(coding)
    update_examples(wb["00A_填写示例"], annotator)
    update_field_guide(wb["00B_字段速查"])
    update_summary_formulas(wb["05_样本总结"], summary_rows, coding)
    update_version_text(wb, annotator)
    wb.save(output_path)
    return {
        "annotator": annotator,
        "input_path": str(input_path),
        "input_sha256": sha256(input_path),
        "output_path": str(output_path),
        "output_sha256": sha256(output_path),
        "standard_sample_summary_rows": len(summary_rows),
    }


def main() -> None:
    if len(sys.argv) != 4:
        raise SystemExit("Usage: remove_pilot0_event_gap_v2_3_3.py <v2.3.2-A.xlsx> <v2.3.2-B.xlsx> <output-dir>")
    input_a = Path(sys.argv[1]).resolve()
    input_b = Path(sys.argv[2]).resolve()
    output_dir = Path(sys.argv[3]).resolve()
    output_dir.mkdir(parents=True, exist_ok=True)

    outputs = []
    for annotator, input_path in (("A", input_a), ("B", input_b)):
        output_path = output_dir / f"ClassIn_IM_Pilot0_可溯源开放编码_v2.3.3_标注员{annotator}_事项缺口已删除受限版.xlsx"
        outputs.append(transform(input_path, output_path, annotator))

    manifest = {
        "status": "TRACEABLE_OPEN_CODING_V2_3_3_EVENT_GAP_REMOVED_PENDING_QA",
        "source_version": "v2.3.2",
        "version": "v2.3.3",
        "review_state": "METHOD_AND_USABILITY_REVIEW_BEFORE_REAL_IM_CONTEXT_ANNOTATION",
        "changes": {
            "removed_field": "事项关键缺口/备注（事项层）",
            "state_rule": "新请求尚未出现结果时只标事项可见状态=未看到，不创造缺口",
            "evidence_limit_rule": "窗口截断、附件未核验、关键前文缺失等统一写在消息层证据边界/推断警报",
            "layout": "03 样本层从 AH:AR 左移到 AG:AQ；05 只读镜像公式同步重写",
        },
        "preservation_policy": "03 A:AF 的原始证据与既有标注值保持不变；原 AG 事项缺口列被物理删除",
        "builder_script_sha256": sha256(Path(__file__)),
        "outputs": outputs,
    }
    manifest_path = output_dir / "pilot0_event_gap_removed_manifest_v2_3_3.json"
    manifest_path.write_text(json.dumps(manifest, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    manifest["manifest_path"] = str(manifest_path)
    print(json.dumps(manifest, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
