#!/usr/bin/env python3
"""Create Topic Pilot 50 v1.2 from the completed human v1 workbook.

The source workbook remains immutable. The derived workbook corrects four
confirmed annotation issues, removes the ambiguous "无教学诉求" suffix,
standardises visible-help wording, and keeps J as the sole Topic column.
"""

from __future__ import annotations

import hashlib
import json
import re
import sys
from datetime import datetime
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side


NAVY = "17384B"
WHITE = "FFFFFF"
GRAY2 = "F4F6F7"
YELLOW = "FFF2CC"
BORDER_COLOR = "B8C4CC"
THIN = Side(style="thin", color=BORDER_COLOR)
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

CORRECTED_TOPICS = {
    "TP50-B3-DIR-04": (
        "T1. 持续讨论电影角色与CP立场、剧情反转、影片彩蛋及一刷/二刷感受\n"
        "T2. 提及开学和课外班即将开始，数学、物理、英语及小四门作业尚未完成，表达时间压力，未见明确求助"
    ),
    "TP50-B3-DIR-05": (
        "T1. 学生持续询问薄膜材料制备、碲化镉回收和应用、压电发电与压力传感器，老师连续解答\n"
        "T2. 围绕中心对称性、电荷分离、极化与晶体结构进行进一步知识讲解\n"
        "T3. 老师反馈学生试卷完成良好、成绩已审核提交，学生询问错题与成绩查看\n"
        "T4. 师生交流老师转行和申请材料学博士的计划"
    ),
    "TP50-B3-DIR-06": (
        "T1. 讨论海外教师当月收入、收入目标、学生留存与教师招聘困难\n"
        "T2. 讨论Viber/WhatsApp的沟通方式、验证码与WhatsApp账号频繁被封问题\n"
        "T3. 安排观课，并讨论通过实物教具提高课堂趣味性和不同教师的教学方法\n"
        "T4. 讨论海外教师合同、护照登记、突然停课后的责任以及跨国追责的实际成本"
    ),
}

AMBIGUOUS_PHRASES = (
    "无教学诉求",
    "无明确教学诉求",
    "无具体教学内容讨论",
)


def sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def value_hash(values: list[Any]) -> str:
    payload = "\n\x1e\n".join("" if value is None else str(value) for value in values)
    return hashlib.sha256(payload.encode("utf-8")).hexdigest()


def topic_lines(value: Any) -> list[str]:
    return [line.strip() for line in str(value or "").splitlines() if line.strip()]


def renumber(lines: list[str]) -> str:
    cleaned = [re.sub(r"^T\d+\.\s*", "", line) for line in lines]
    return "\n".join(f"T{index}. {line}" for index, line in enumerate(cleaned, start=1))


def calibrate_wording(value: str) -> str:
    calibrated = value
    # Remove the ambiguous negative classification from neutral Topic names.
    for phrase in AMBIGUOUS_PHRASES:
        calibrated = calibrated.replace(f"，{phrase}", "").replace(phrase, "")
    # Describe only what is observable in the current 100-message window.
    calibrated = calibrated.replace("无进一步求助", "未见进一步求助")
    calibrated = calibrated.replace("未提出求助", "未见明确求助")
    calibrated = re.sub(r"，{2,}", "，", calibrated)
    return calibrated.strip("， ")


def style_instruction_row(ws: Any, row: int, label: str, explanation: str) -> None:
    ws.cell(row, 1, label)
    ws.cell(row, 1).fill = PatternFill("solid", fgColor=NAVY)
    ws.cell(row, 1).font = Font(color=WHITE, bold=True)
    ws.cell(row, 1).alignment = Alignment(vertical="top", wrap_text=True)
    for merged in list(ws.merged_cells.ranges):
        if merged.min_row == row and merged.max_row == row:
            ws.unmerge_cells(str(merged))
    ws.merge_cells(start_row=row, start_column=2, end_row=row, end_column=10)
    ws.cell(row, 2, explanation)
    ws.cell(row, 2).fill = PatternFill("solid", fgColor=GRAY2)
    ws.cell(row, 2).alignment = Alignment(vertical="top", wrap_text=True)
    for col in range(1, 11):
        ws.cell(row, col).border = BORDER
    ws.row_dimensions[row].height = 60


def main() -> None:
    if len(sys.argv) != 4:
        raise SystemExit("Usage: calibrate_topic_pilot50_v1_2.py <completed-v1.xlsx> <v1.2.xlsx> <manifest.json>")
    source_path, output_path, manifest_path = map(lambda item: Path(item).resolve(), sys.argv[1:])
    wb = load_workbook(source_path)
    expected_sheets = [
        "00_使用说明", "01_案例目录", "02_案例消息", "03_Topic人工标注",
        "04_问题与校准", "05_抽样审计", "06_原表字段索引",
    ]
    if wb.sheetnames != expected_sheets:
        raise RuntimeError("Unexpected completed-v1 workbook structure")

    topic_ws = wb["03_Topic人工标注"]
    source_topics = [topic_ws.cell(row, 10).value for row in range(5, 55)]
    if sum(bool(value) for value in source_topics) != 50:
        raise RuntimeError("Expected all 50 Topic cells to be filled")
    if any(topic_ws.cell(row, col).value not in (None, "") for row in range(5, 55) for col in (11, 12)):
        raise RuntimeError("K/L contain unexpected human input")

    source_ambiguous_phrase_counts = {
        phrase: sum(str(value or "").count(phrase) for value in source_topics)
        for phrase in AMBIGUOUS_PHRASES
    }
    source_help_wording_counts = {
        phrase: sum(str(value or "").count(phrase) for value in source_topics)
        for phrase in ("无进一步求助", "未提出求助")
    }

    row_by_case = {topic_ws.cell(row, 2).value: row for row in range(5, 55)}
    changes: list[dict[str, Any]] = []

    # Correct the three confirmed cross-case Topic placements.
    for case_id, corrected in CORRECTED_TOPICS.items():
        row = row_by_case[case_id]
        before = str(topic_ws.cell(row, 10).value or "")
        topic_ws.cell(row, 10, corrected)
        changes.append({"case_id": case_id, "change": "replace_cross_case_topic", "before": before, "after": corrected})

    # Remove the under-threshold WeChat/QQ Topic from B1-GRP-02.
    boundary_case = "TP50-B1-GRP-02"
    boundary_row = row_by_case[boundary_case]
    before = str(topic_ws.cell(boundary_row, 10).value or "")
    kept = [line for line in topic_lines(before) if "微信" not in line and "QQ" not in line]
    after = renumber(kept)
    topic_ws.cell(boundary_row, 10, after)
    changes.append({"case_id": boundary_case, "change": "remove_under_threshold_topic", "before": before, "after": after})

    # Calibrate all wording without changing the underlying topic fact.
    wording_changed_cells = 0
    removed_phrase_counts = {phrase: 0 for phrase in AMBIGUOUS_PHRASES}
    help_wording_counts = {"无进一步求助": 0, "未提出求助": 0}
    for row in range(5, 55):
        current = str(topic_ws.cell(row, 10).value or "")
        for phrase in AMBIGUOUS_PHRASES:
            removed_phrase_counts[phrase] += current.count(phrase)
        for phrase in help_wording_counts:
            help_wording_counts[phrase] += current.count(phrase)
        calibrated = calibrate_wording(current)
        if calibrated != current:
            wording_changed_cells += 1
            topic_ws.cell(row, 10, calibrated)

    calibrated_topics = [topic_ws.cell(row, 10).value for row in range(5, 55)]
    if any(phrase in str(value) for value in calibrated_topics for phrase in AMBIGUOUS_PHRASES):
        raise RuntimeError("Ambiguous teaching-request wording remains")

    # Keep J as the sole Topic column.
    for merged in ("A1:M1", "A2:M2"):
        if merged in {str(item) for item in topic_ws.merged_cells.ranges}:
            topic_ws.unmerge_cells(merged)
    topic_ws.delete_cols(11, 3)
    topic_ws.merge_cells("A1:J1")
    topic_ws.merge_cells("A2:J2")
    topic_ws["A1"] = "Topic Pilot 50 v1.2 · 人工 Topic 校准版"
    topic_ws["A2"] = "J列为已完成并校准的Topic；已修正3处串位、1个准入边界，并移除含混的“无教学诉求”。"
    topic_ws["J4"] = "校准后Topic【中性事实描述；教学关联与求助后置分维度判断】"
    topic_ws.auto_filter.ref = "A4:J54"
    topic_ws.freeze_panes = "J5"
    topic_ws.data_validations.dataValidation = []
    topic_ws.column_dimensions["J"].width = 76
    for row in range(5, 55):
        topic_ws.cell(row, 10).fill = PatternFill("solid", fgColor=YELLOW)
        topic_ws.cell(row, 10).alignment = Alignment(vertical="top", wrap_text=True)

    guide_ws = wb["00_使用说明"]
    guide_ws["A1"] = "ClassIn IM Topic Pilot 50 v1.2 · Topic 校准版"
    guide_ws["A2"] = "当前J列是人工Topic的校准派生层；原v1人工填写保留不覆写。"
    style_instruction_row(
        guide_ws, 4, "Topic写法",
        "Topic只写中性可见事实，不在名称末尾追加“无教学诉求”。",
    )
    style_instruction_row(
        guide_ws, 7, "为什么移除否定句",
        "学习内容、明确求助、知识教学、教学服务和教学管理是不同维度，不能用一句“无教学诉求”同时否定。",
    )
    style_instruction_row(
        guide_ws, 8, "可见求助表轰",
        "只在100条窗口有证据时写“未见明确求助”；不把“没有求助”扩大为“与教学无关”。",
    )
    style_instruction_row(
        guide_ws, 12, "后置拆解",
        "AI后续独立生成：内容领域、沟通目的、教学关联类型、是否明确求助、证据回链、有效发言数和置信度。",
    )

    calibration_ws = wb["04_问题与校准"]
    calibration_ws["A1"] = "Topic Pilot 50 v1.2 · 校准变更记录"
    records = [
        ("CAL-001", "第三批", "TP50-B3-DIR-04", "Topic串入DIR-08家庭争吵", "替换为电影讨论+作业压力", "根据本Case 100条原文修正", "已确认"),
        ("CAL-002", "第三批", "TP50-B3-DIR-05", "Topic串入DIR-09图书馆/作业", "替换为材料科学答疑+试卷反馈", "根据本Case 100条原文修正", "已确认"),
        ("CAL-003", "第三批", "TP50-B3-DIR-06", "Topic串入DIR-10情感交流", "替换为教师运营/工具/观课/合同", "根据本Case 100条原文修正", "已确认"),
        ("CAL-004", "第一批", "TP50-B1-GRP-02", "微信/QQ扩列只有3条有效发言", "删除该Topic并重新编号", "执行同一语义主题≥5条有效发言", "已确认"),
        ("CAL-005", "全部", "ALL", "“无教学诉求”混合多个维度", "从Topic名称移除；求助表轰改为“未见…”", "教学关联/求助/服务/管理后置分维度判断", "已确认"),
    ]
    for row in range(5, 55):
        for col in range(1, 7):
            calibration_ws.cell(row, col, None)
        calibration_ws.cell(row, 7, None)
    for row, record in enumerate(records, start=5):
        for col, value in enumerate(record, start=1):
            calibration_ws.cell(row, col, value)
            calibration_ws.cell(row, col).alignment = Alignment(vertical="top", wrap_text=True)
            calibration_ws.cell(row, col).border = BORDER

    audit_ws = wb["05_抽样审计"]
    audit_ws["A16"] = "v1.2校准"
    audit_ws["B16"] = "3处Topic串位已修正；1个低于5条有效发言的Topic已删除；“无教学诉求”已从Topic名称移除"
    for col in (1, 2):
        audit_ws.cell(16, col).fill = PatternFill("solid", fgColor=GRAY2)
        audit_ws.cell(16, col).alignment = Alignment(vertical="top", wrap_text=True)
        audit_ws.cell(16, col).border = BORDER

    wb.active = wb.sheetnames.index("03_Topic人工标注")
    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)

    manifest = {
        "status": "TOPIC_PILOT50_V1_2_CALIBRATED_READY_FOR_AI_EVIDENCE_MAPPING",
        "version": "v1.2",
        "generated_at": datetime.now().isoformat(timespec="seconds"),
        "source_workbook": {"path": str(source_path), "sha256": sha256(source_path)},
        "output_workbook": {"path": str(output_path), "sha256": sha256(output_path)},
        "calibration": {
            "confirmed_case_changes": [item["case_id"] for item in changes],
            "change_count": len(changes),
            "source_ambiguous_phrase_counts": source_ambiguous_phrase_counts,
            "source_help_wording_counts": source_help_wording_counts,
            "post_correction_wording_removals": removed_phrase_counts,
            "post_correction_help_wording_standardisations": help_wording_counts,
            "changed_topic_cells_for_wording": wording_changed_cells,
            "source_topic_lines": sum(len(topic_lines(value)) for value in source_topics),
            "output_topic_lines": sum(len(topic_lines(value)) for value in calibrated_topics),
            "source_topic_values_sha256": value_hash(source_topics),
            "output_topic_values_sha256": value_hash(calibrated_topics),
        },
        "preservation": {
            "source_not_overwritten": True,
            "filled_topic_cells": sum(bool(value) for value in calibrated_topics),
            "sole_topic_column": "03_Topic人工标注!J5:J54",
            "evidence_sheets_preserved": ["01_案例目录", "02_案例消息", "06_原表字段索引"],
        },
        "builder": {"path": str(Path(__file__).resolve()), "sha256": sha256(Path(__file__).resolve())},
    }
    manifest_path.write_text(json.dumps(manifest, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps(manifest, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
