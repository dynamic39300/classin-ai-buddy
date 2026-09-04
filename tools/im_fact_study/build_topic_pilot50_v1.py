#!/usr/bin/env python3
"""Build the ClassIn IM Topic Pilot 50 human-annotation package.

The only research input is the locked original Excel workbook. Sampling is
content-blind and uses only mechanical structure: chat type, sender count,
role pattern, group-size band, message completeness, and visible time span.
No prior semantic labels, topic conclusions, or derived sample workbooks are
read by this script.
"""

from __future__ import annotations

import hashlib
import json
import random
import sys
from collections import Counter, defaultdict
from datetime import datetime
from pathlib import Path
from typing import Any, Iterable

from openpyxl import Workbook, load_workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation


VERSION = "v1"
SEED = 20260830
SOURCE_SHEET = "Sheet"
EXPECTED_SOURCE_ROWS = 819_700
EXPECTED_MESSAGES_PER_CASE = 100
# This case was already exposed to the human annotator in the retired
# per-message trial. Excluding it protects the blind-validation batch and is
# not a semantic sampling input.
EXCLUDED_CLUSTER_IDS = {"27324211"}

COL = {
    "identity": 0,
    "user_type": 1,
    "user_num": 2,
    "id": 3,
    "clusterid": 4,
    "clustertype": 5,
    "msgbucketid": 6,
    "msgid": 7,
    "msgcmd": 8,
    "msgdata": 9,
    "replymsgid": 10,
    "sourceuid": 11,
    "targetuids": 12,
    "timetag": 13,
    "timeformat": 14,
    "dt": 15,
    "rn": 16,
    "content": 17,
    "strtalker": 18,
    "readable_time": 19,
}

COLORS = {
    "navy": "17384B",
    "blue": "DCEEF7",
    "gray": "E9EEF1",
    "gray2": "F4F6F7",
    "yellow": "FFF2CC",
    "green": "E2F0D9",
    "orange": "FCE4D6",
    "purple": "E4DFEC",
    "white": "FFFFFF",
    "ink": "1C2B33",
    "border": "B8C4CC",
}

THIN = Side(style="thin", color=COLORS["border"])
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def safe_text(value: Any) -> str:
    if value is None:
        return ""
    return ILLEGAL_CHARACTERS_RE.sub("", str(value))


def parse_time(value: Any) -> datetime | None:
    if isinstance(value, datetime):
        return value
    text = safe_text(value).strip()
    if not text:
        return None
    try:
        return datetime.fromisoformat(text)
    except ValueError:
        return None


def sender_bucket(count: int) -> str:
    if count <= 1:
        return "1"
    if count == 2:
        return "2"
    if count <= 5:
        return "3-5"
    return "6+"


def duration_bucket(days: float) -> str:
    if days <= 1:
        return "≤1天"
    if days <= 7:
        return "2-7天"
    return ">7天"


def role_bucket(roles: set[str]) -> str:
    has_student = bool(roles & {"学生", "旁听生"})
    has_staff = bool(roles & {"教师", "班主任"})
    if has_student and has_staff:
        return "师生可见"
    if has_student:
        return "仅学生角色可见"
    if has_staff:
        return "仅教职角色可见"
    return "仅其他角色可见"


def structural_stratum(info: dict[str, Any]) -> str:
    common = f"发送者{sender_bucket(info['sender_count'])}|跨度{duration_bucket(info['span_days'])}"
    if info["clustertype"] == 1:
        return common
    return f"{common}|{role_bucket(info['roles'])}|{info['group_size']}"


def finalize_info(info: dict[str, Any]) -> None:
    info["sender_count"] = len(info["senders"])
    info["role_pattern"] = "+".join(sorted(info["roles"])) or "未见"
    info["group_size"] = "+".join(sorted(info["group_sizes"])) or "未见"
    info["text_ratio"] = info["text_count"] / info["message_count"] if info["message_count"] else 0
    if info["first_time"] and info["last_time"]:
        info["span_days"] = max(0.0, (info["last_time"] - info["first_time"]).total_seconds() / 86400)
    else:
        info["span_days"] = 0.0
    info["stratum"] = structural_stratum(info)


def scan_source(source_path: Path) -> tuple[list[str], dict[tuple[int, str], dict[str, Any]], int]:
    wb = load_workbook(source_path, read_only=True, data_only=True)
    ws = wb[SOURCE_SHEET]
    header = [safe_text(v) for v in next(ws.iter_rows(min_row=1, max_row=1, values_only=True))]
    clusters: dict[tuple[int, str], dict[str, Any]] = {}
    row_count = 0
    for excel_row, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        row_count += 1
        chat_type = int(row[COL["clustertype"]])
        clusterid = safe_text(row[COL["clusterid"]])
        key = (chat_type, clusterid)
        current_time = parse_time(row[COL["readable_time"]]) or parse_time(row[COL["timeformat"]])
        info = clusters.get(key)
        if info is None:
            info = {
                "clustertype": chat_type,
                "clusterid": clusterid,
                "message_count": 0,
                "senders": set(),
                "roles": set(),
                "group_sizes": set(),
                "text_count": 0,
                "first_time": current_time,
                "last_time": current_time,
                "first_excel_row": excel_row,
                "last_excel_row": excel_row,
            }
            clusters[key] = info
        info["message_count"] += 1
        info["senders"].add(safe_text(row[COL["sourceuid"]]))
        info["roles"].add(safe_text(row[COL["user_type"]]) or "未见")
        info["group_sizes"].add(safe_text(row[COL["user_num"]]) or "未见")
        if safe_text(row[COL["content"]]).strip():
            info["text_count"] += 1
        if current_time is not None:
            if info["first_time"] is None or current_time < info["first_time"]:
                info["first_time"] = current_time
            if info["last_time"] is None or current_time > info["last_time"]:
                info["last_time"] = current_time
        info["last_excel_row"] = excel_row
    wb.close()
    for info in clusters.values():
        finalize_info(info)
    return header, clusters, row_count


def diversified_sample(candidates: Iterable[dict[str, Any]], count: int, rng: random.Random) -> list[dict[str, Any]]:
    grouped: dict[str, list[dict[str, Any]]] = defaultdict(list)
    for item in candidates:
        grouped[item["stratum"]].append(item)
    for items in grouped.values():
        rng.shuffle(items)
    strata = list(grouped)
    rng.shuffle(strata)
    selected: list[dict[str, Any]] = []
    while len(selected) < count:
        progressed = False
        for stratum in strata:
            if grouped[stratum]:
                selected.append(grouped[stratum].pop())
                progressed = True
                if len(selected) == count:
                    break
        if not progressed:
            raise RuntimeError(f"Not enough candidates: requested {count}, selected {len(selected)}")
        rng.shuffle(strata)
    return selected


def quota_sample(
    candidates: Iterable[dict[str, Any]],
    quotas: dict[str, int],
    rng: random.Random,
) -> list[dict[str, Any]]:
    """Sample explicit sender-count buckets, diversified by other structure."""
    pool = list(candidates)
    picks: list[dict[str, Any]] = []
    for bucket, count in quotas.items():
        bucket_pool = [x for x in pool if sender_bucket(x["sender_count"]) == bucket]
        if len(bucket_pool) < count:
            raise RuntimeError(f"Not enough candidates in sender bucket {bucket}: {len(bucket_pool)} < {count}")
        picks.extend(diversified_sample(bucket_pool, count, rng))
    return picks


def choose_cases(clusters: dict[tuple[int, str], dict[str, Any]]) -> list[dict[str, Any]]:
    rng = random.Random(SEED)
    eligible = [
        x for x in clusters.values()
        if x["message_count"] == EXPECTED_MESSAGES_PER_CASE
        and x["clusterid"] not in EXCLUDED_CLUSTER_IDS
    ]
    selected: list[dict[str, Any]] = []
    used: set[tuple[int, str]] = set()

    direct_quotas = {
        1: {"2": 5},
        2: {"1": 2, "2": 8},
        3: {"1": 1, "2": 9},
    }
    group_quotas = {
        2: {"1": 1, "2": 2, "3-5": 3, "6+": 4},
        3: {"1": 1, "2": 2, "3-5": 3, "6+": 4},
    }

    direct_pool = [x for x in eligible if x["clustertype"] == 1]
    for batch, quotas in direct_quotas.items():
        remaining = [x for x in direct_pool if (1, x["clusterid"]) not in used]
        if batch == 1:
            preferred = [x for x in remaining if x["text_ratio"] >= 0.95 and x["span_days"] <= 7]
            source = preferred if len(preferred) >= sum(quotas.values()) else remaining
        else:
            source = remaining
        for item in quota_sample(source, quotas, rng):
            item = dict(item)
            item["batch"] = batch
            selected.append(item)
            used.add((1, item["clusterid"]))

    group_pool = [x for x in eligible if x["clustertype"] == 0]
    preferred = [
        x for x in group_pool
        if x["sender_count"] >= 3
        and role_bucket(x["roles"]) == "师生可见"
        and x["text_ratio"] >= 0.95
        and x["span_days"] <= 7
    ]
    for item in diversified_sample(preferred, 5, rng):
        item = dict(item)
        item["batch"] = 1
        selected.append(item)
        used.add((0, item["clusterid"]))

    for batch, quotas in group_quotas.items():
        remaining = [x for x in group_pool if (0, x["clusterid"]) not in used]
        for item in quota_sample(remaining, quotas, rng):
            item = dict(item)
            item["batch"] = batch
            selected.append(item)
            used.add((0, item["clusterid"]))

    batch_labels = {
        1: "第一批｜方法探索（10例）",
        2: "第二批｜规则磨合（20例）",
        3: "第三批｜盲测验证（20例）",
    }
    for batch in (1, 2, 3):
        for chat_type, prefix in ((1, "DIR"), (0, "GRP")):
            items = [x for x in selected if x["batch"] == batch and x["clustertype"] == chat_type]
            items.sort(key=lambda x: hashlib.sha256(f"{SEED}|{x['clusterid']}".encode()).hexdigest())
            for index, item in enumerate(items, start=1):
                item["case_id"] = f"TP50-B{batch}-{prefix}-{index:02d}"
                item["batch_label"] = batch_labels[batch]
                item["chat_type_label"] = "1v1单聊" if chat_type == 1 else "班级群聊"
    selected.sort(key=lambda x: (x["batch"], x["clustertype"] != 1, x["case_id"]))
    return selected


def extract_selected_rows(source_path: Path, selected: list[dict[str, Any]]) -> dict[tuple[int, str], list[tuple[int, tuple[Any, ...]]]]:
    selected_keys = {(x["clustertype"], x["clusterid"]) for x in selected}
    rows: dict[tuple[int, str], list[tuple[int, tuple[Any, ...]]]] = defaultdict(list)
    wb = load_workbook(source_path, read_only=True, data_only=True)
    ws = wb[SOURCE_SHEET]
    for excel_row, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        key = (int(row[COL["clustertype"]]), safe_text(row[COL["clusterid"]]))
        if key in selected_keys:
            rows[key].append((excel_row, row))
    wb.close()
    for item in selected:
        key = (item["clustertype"], item["clusterid"])
        if len(rows[key]) != EXPECTED_MESSAGES_PER_CASE:
            raise RuntimeError(f"{key} extracted {len(rows[key])} rows, expected {EXPECTED_MESSAGES_PER_CASE}")
    return rows


def style_title(ws: Any, title: str, subtitle: str, end_col: int) -> None:
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=end_col)
    ws.cell(1, 1, title)
    ws.cell(1, 1).fill = PatternFill("solid", fgColor=COLORS["navy"])
    ws.cell(1, 1).font = Font(color=COLORS["white"], bold=True, size=16)
    ws.cell(1, 1).alignment = Alignment(vertical="center")
    ws.row_dimensions[1].height = 30
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=end_col)
    ws.cell(2, 1, subtitle)
    ws.cell(2, 1).fill = PatternFill("solid", fgColor=COLORS["blue"])
    ws.cell(2, 1).font = Font(color=COLORS["ink"], size=10)
    ws.cell(2, 1).alignment = Alignment(wrap_text=True, vertical="center")
    ws.row_dimensions[2].height = 34


def style_header(ws: Any, row: int, start_col: int, end_col: int, fill: str | None = None) -> None:
    for col in range(start_col, end_col + 1):
        cell = ws.cell(row, col)
        cell.fill = PatternFill("solid", fgColor=fill or COLORS["navy"])
        cell.font = Font(color=COLORS["white"], bold=True, size=9)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = BORDER


def style_body(ws: Any, min_row: int, max_row: int, min_col: int, max_col: int, fill: str = "FFFFFF") -> None:
    for row in ws.iter_rows(min_row=min_row, max_row=max_row, min_col=min_col, max_col=max_col):
        for cell in row:
            cell.fill = PatternFill("solid", fgColor=fill)
            cell.font = Font(color=COLORS["ink"], size=9)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = BORDER


def set_widths(ws: Any, widths: dict[str, float]) -> None:
    for column, width in widths.items():
        ws.column_dimensions[column].width = width


def add_status_validation(ws: Any, cells: str) -> None:
    values = "未开始,进行中,已完成,需讨论"
    dv = DataValidation(type="list", formula1=f'"{values}"', allow_blank=False)
    dv.promptTitle = "选择标注状态"
    dv.prompt = values.replace(",", " / ")
    dv.errorTitle = "请使用规定取值"
    dv.error = values.replace(",", " / ")
    dv.errorStyle = "stop"
    dv.showErrorMessage = True
    dv.showInputMessage = True
    ws.add_data_validation(dv)
    dv.add(cells)


def case_digest(case_rows: list[tuple[int, tuple[Any, ...]]]) -> str:
    digest = hashlib.sha256()
    for excel_row, row in case_rows:
        fields = [
            excel_row,
            row[COL["clusterid"]],
            row[COL["id"]],
            row[COL["msgid"]],
            row[COL["sourceuid"]],
            row[COL["readable_time"]],
            row[COL["content"]],
        ]
        digest.update(("\x1f".join(safe_text(x) for x in fields) + "\n").encode("utf-8"))
    return digest.hexdigest()


def build_workbook(
    source_path: Path,
    source_sha: str,
    source_header: list[str],
    source_row_count: int,
    clusters: dict[tuple[int, str], dict[str, Any]],
    selected: list[dict[str, Any]],
    selected_rows: dict[tuple[int, str], list[tuple[int, tuple[Any, ...]]]],
    output_path: Path,
) -> dict[str, Any]:
    wb = Workbook()
    wb.remove(wb.active)

    # 00 Instructions
    ws = wb.create_sheet("00_使用说明")
    style_title(
        ws,
        "ClassIn IM Topic Pilot 50 · 人工 Topic 标注",
        "只使用原始 Excel 重新抽样。人工以100条消息案例为单位提炼 Topic；单条消息只作为关键证据，不再逐句编码。",
        10,
    )
    instruction_rows = [
        (4, "你只需要填写哪里", "03_Topic人工标注 中黄色的4列：人工Topic初稿、关键证据、上下文不足/不确定、标注状态。其他内容只读。"),
        (5, "第一步", "按批次开始：第一批10例 → 第二批20例 → 第三批20例。先不要跳到第三批。"),
        (6, "第二步", "点击‘打开100条消息’，快速通读同一 case_id 的全部消息。"),
        (7, "第三步", "在人工Topic初稿中一行写一个 Topic，可使用 T1/T2/T3…；Topic 数量完全按实际内容，不设上限。"),
        (8, "第四步", "建议写关键证据的 case_seq、msgid 或 rn；不确定时保留不确定，不猜角色、动机、结果或缺失对话。"),
        (9, "Topic写法", "Good：‘讨论课程时间冲突，并提出请假请求’。Bad：‘需要AI请假助手’（产品方案不是聊天事实）。"),
        (10, "表达不等于诉求", "‘今天的题太难了’可记录为学习感受，但不能自动写成‘寻求题目讲解’；只有可见求助证据时才写诉求。"),
        (11, "闲聊也是有效结果", "明星、天气、吃饭、家庭琐事等可记录为闲聊 Topic；它们帮助我们识别哪些内容不应被产品化。"),
        (12, "多Topic", "一个案例有多个主题就分行列出；同一沟通片段的连续推进不要机械地一条消息拆一个 Topic。"),
        (13, "第三批含义", "第三批是盲测验证：前两批规则稳定后再标，用于检查方法是否能迁移，不应在标注前查看AI拆解。"),
    ]
    for row, label, explanation in instruction_rows:
        ws.cell(row, 1, label)
        ws.cell(row, 1).fill = PatternFill("solid", fgColor=COLORS["navy"])
        ws.cell(row, 1).font = Font(color=COLORS["white"], bold=True)
        ws.cell(row, 1).alignment = Alignment(vertical="top", wrap_text=True)
        ws.merge_cells(start_row=row, start_column=2, end_row=row, end_column=10)
        ws.cell(row, 2, explanation)
        ws.cell(row, 2).fill = PatternFill("solid", fgColor=COLORS["gray2"])
        ws.cell(row, 2).alignment = Alignment(vertical="top", wrap_text=True)
        for col in range(1, 11):
            ws.cell(row, col).border = BORDER
        ws.row_dimensions[row].height = 42
    ws["A15"] = "人工Topic初稿示例"
    ws["A15"].fill = PatternFill("solid", fgColor=COLORS["navy"])
    ws["A15"].font = Font(color=COLORS["white"], bold=True)
    ws.merge_cells("B15:J15")
    ws["B15"] = (
        "T1｜讨论周末聚餐安排，无教学诉求｜证据 case_seq 12-24\n"
        "T2｜询问课程开始时间，并提出调课请求｜证据 case_seq 55-63\n"
        "T3｜反馈作业文件无法打开，请求重新发送｜证据 msgid 12345/12346"
    )
    ws["B15"].fill = PatternFill("solid", fgColor=COLORS["yellow"])
    ws["B15"].alignment = Alignment(wrap_text=True, vertical="top")
    for col in range(1, 11):
        ws.cell(15, col).border = BORDER
    ws.row_dimensions[15].height = 78
    set_widths(ws, {"A": 24, "B": 22, "C": 18, "D": 18, "E": 18, "F": 18, "G": 18, "H": 18, "I": 18, "J": 18})
    ws.sheet_view.showGridLines = False

    # 01 Case catalog
    ws = wb.create_sheet("01_案例目录")
    style_title(ws, "Topic Pilot 50 案例目录", "50个案例已一次性抽完；按批次和 case_id 排序。抽样只使用机械结构，不读取消息关键词或既有主题。", 17)
    headers = [
        "批次", "case_id", "聊天类型", "clusterid", "消息数", "可见发送者数", "发送者结构", "可见角色",
        "群规模字段", "开始时间", "结束时间", "跨度天数", "文本完整率", "机械分层", "原表最小-最大行（可能非连续）", "案例证据SHA256", "打开消息",
    ]
    ws.append([])
    ws.append(headers)
    style_header(ws, 4, 1, len(headers))
    ws.row_dimensions[4].height = 54

    # Raw message sheet first, to know hyperlink rows.
    msg_ws = wb.create_sheet("02_案例消息")
    style_title(msg_ws, "Topic Pilot 50 原始消息证据", "每个 case_id 恰好100条原表消息。建议使用筛选器选择 case_id；A:I 为主要阅读区，J:T 为原表溯源字段。整页只读。", 20)
    msg_headers = [
        "case_id", "批次", "聊天类型", "case_seq", "原表时间", "strtalker", "user_type", "sourceuid", "原表正文",
        "原表Excel行", "clusterid", "clustertype", "id", "msgid", "replymsgid", "rn", "identity", "user_num", "targetuids", "msgdata",
    ]
    msg_ws.append([])
    msg_ws.append(msg_headers)
    style_header(msg_ws, 4, 1, len(msg_headers))
    msg_ws.row_dimensions[4].height = 58
    message_start_rows: dict[str, int] = {}
    message_end_rows: dict[str, int] = {}
    current_row = 5
    for item in selected:
        key = (item["clustertype"], item["clusterid"])
        case_rows = selected_rows[key]
        message_start_rows[item["case_id"]] = current_row
        for seq, (excel_row, row) in enumerate(case_rows, start=1):
            values = [
                item["case_id"], item["batch_label"], item["chat_type_label"], seq,
                safe_text(row[COL["readable_time"]]), safe_text(row[COL["strtalker"]]),
                safe_text(row[COL["user_type"]]), safe_text(row[COL["sourceuid"]]), safe_text(row[COL["content"]]),
                excel_row, safe_text(row[COL["clusterid"]]), int(row[COL["clustertype"]]), safe_text(row[COL["id"]]),
                safe_text(row[COL["msgid"]]), safe_text(row[COL["replymsgid"]]), row[COL["rn"]], safe_text(row[COL["identity"]]),
                safe_text(row[COL["user_num"]]), safe_text(row[COL["targetuids"]]), safe_text(row[COL["msgdata"]]),
            ]
            msg_ws.append(values)
            fill = COLORS["blue"] if item["batch"] == 1 else COLORS["green"] if item["batch"] == 2 else COLORS["purple"]
            for col in range(1, 21):
                cell = msg_ws.cell(current_row, col)
                cell.fill = PatternFill("solid", fgColor=fill if col <= 4 else COLORS["gray2"])
                cell.font = Font(color=COLORS["ink"], size=9)
                cell.alignment = Alignment(vertical="top", wrap_text=True)
                cell.border = BORDER
            if seq == 1:
                for col in range(1, 21):
                    msg_ws.cell(current_row, col).border = Border(left=THIN, right=THIN, top=Side(style="medium", color=COLORS["navy"]), bottom=THIN)
            current_row += 1
        message_end_rows[item["case_id"]] = current_row - 1
    msg_ws.auto_filter.ref = f"A4:T{msg_ws.max_row}"
    msg_ws.freeze_panes = "E5"
    msg_ws.sheet_view.showGridLines = False
    set_widths(msg_ws, {
        "A": 20, "B": 25, "C": 12, "D": 10, "E": 20, "F": 18, "G": 14, "H": 16, "I": 64,
        "J": 14, "K": 15, "L": 12, "M": 16, "N": 14, "O": 14, "P": 9, "Q": 11, "R": 24, "S": 16, "T": 42,
    })
    for col in ("J", "K", "L", "M", "N", "O", "P", "Q", "R", "S", "T"):
        msg_ws.column_dimensions[col].hidden = True

    # Complete case catalog now that target rows are known.
    for item in selected:
        key = (item["clustertype"], item["clusterid"])
        item["case_sha256"] = case_digest(selected_rows[key])
        first_time = item["first_time"].isoformat(sep=" ") if item["first_time"] else ""
        last_time = item["last_time"].isoformat(sep=" ") if item["last_time"] else ""
        row = [
            item["batch_label"], item["case_id"], item["chat_type_label"], item["clusterid"], item["message_count"],
            item["sender_count"], sender_bucket(item["sender_count"]), item["role_pattern"], item["group_size"],
            first_time, last_time, round(item["span_days"], 2), round(item["text_ratio"], 4), item["stratum"],
            f"{item['first_excel_row']}-{item['last_excel_row']}", item["case_sha256"],
            f'=HYPERLINK("#\'02_案例消息\'!A{message_start_rows[item["case_id"]]}","打开100条消息")',
        ]
        ws.append(row)
    style_body(ws, 5, ws.max_row, 1, len(headers), COLORS["gray2"])
    for row in range(5, ws.max_row + 1):
        ws.cell(row, 17).font = Font(color="0563C1", underline="single")
    ws.auto_filter.ref = f"A4:Q{ws.max_row}"
    ws.freeze_panes = "A5"
    ws.sheet_view.showGridLines = False
    set_widths(ws, {"A": 25, "B": 20, "C": 12, "D": 15, "E": 9, "F": 12, "G": 12, "H": 27, "I": 25, "J": 20, "K": 20, "L": 12, "M": 13, "N": 56, "O": 18, "P": 68, "Q": 18})

    # 03 Human topic annotation
    ws = wb.create_sheet("03_Topic人工标注")
    style_title(
        ws,
        "Topic Pilot 50 · 人工 Topic 初稿",
        "只填写黄色列 J:M。每个案例一行；Topic初稿单元格内可用 Alt+Enter 写任意数量的 T1/T2/T3…，不预设Topic数量。",
        13,
    )
    topic_headers = [
        "批次", "case_id", "聊天类型", "clusterid", "消息数", "可见发送者数", "可见角色", "时间范围", "打开100条消息",
        "人工Topic初稿【必填；一行一个Topic】", "关键证据 case_seq/msgid/rn【建议】", "上下文不足/不确定说明【选填】", "标注状态【下拉】",
    ]
    ws.append([])
    ws.append(topic_headers)
    style_header(ws, 4, 1, 9)
    style_header(ws, 4, 10, 13, fill="9C6500")
    ws.row_dimensions[4].height = 86
    for item in selected:
        start = message_start_rows[item["case_id"]]
        values = [
            item["batch_label"], item["case_id"], item["chat_type_label"], item["clusterid"], 100,
            item["sender_count"], item["role_pattern"],
            f"{item['first_time'].isoformat(sep=' ') if item['first_time'] else ''} → {item['last_time'].isoformat(sep=' ') if item['last_time'] else ''}",
            f'=HYPERLINK("#\'02_案例消息\'!A{start}","打开100条消息")', "", "", "", "未开始",
        ]
        ws.append(values)
    style_body(ws, 5, ws.max_row, 1, 9, COLORS["gray2"])
    style_body(ws, 5, ws.max_row, 10, 13, COLORS["yellow"])
    for row in range(5, ws.max_row + 1):
        ws.row_dimensions[row].height = 92
        ws.cell(row, 9).font = Font(color="0563C1", underline="single")
    add_status_validation(ws, f"M5:M{ws.max_row}")
    ws.auto_filter.ref = f"A4:M{ws.max_row}"
    ws.freeze_panes = "J5"
    ws.sheet_view.showGridLines = False
    set_widths(ws, {"A": 25, "B": 20, "C": 12, "D": 15, "E": 9, "F": 12, "G": 27, "H": 38, "I": 18, "J": 62, "K": 38, "L": 42, "M": 14})

    # 04 Questions / calibration notes
    ws = wb.create_sheet("04_问题与校准")
    style_title(ws, "Topic 标注问题与校准记录", "遇到拆分、合并、是否属于诉求或上下文不足等问题时记录在这里；不要为了完成而猜测。", 7)
    headers = ["记录ID", "批次", "case_id", "问题/分歧", "当前临时处理", "最终规则/答复", "状态"]
    ws.append([])
    ws.append(headers)
    style_header(ws, 4, 1, 7)
    for row in range(5, 55):
        ws.append([f"Q{row-4:03d}", "", "", "", "", "", "待讨论"])
    style_body(ws, 5, 54, 1, 7, COLORS["yellow"])
    status = DataValidation(type="list", formula1='"待讨论,已校准,不适用"', allow_blank=False)
    ws.add_data_validation(status)
    status.add("G5:G54")
    ws.freeze_panes = "A5"
    ws.auto_filter.ref = "A4:G54"
    ws.sheet_view.showGridLines = False
    set_widths(ws, {"A": 12, "B": 25, "C": 20, "D": 52, "E": 42, "F": 52, "G": 14})

    # 05 Sampling audit
    ws = wb.create_sheet("05_抽样审计")
    style_title(ws, "Topic Pilot 50 抽样与证据审计", "本页只读。50例从锁定原始Excel直接抽取，固定种子；不使用消息关键词、旧主题标签或旧分析结论。", 8)
    audit_rows = [
        ("原始文件", str(source_path)),
        ("原始文件SHA256", source_sha),
        ("原始工作表", SOURCE_SHEET),
        ("原始消息行数", source_row_count),
        ("原始100消息案例数", len(clusters)),
        ("单聊案例池", sum(1 for x in clusters.values() if x["clustertype"] == 1)),
        ("群聊案例池", sum(1 for x in clusters.values() if x["clustertype"] == 0)),
        ("抽样随机种子", SEED),
        ("人为暴露排除", "clusterid 27324211；曾用于旧方法试标，仅为保护盲测，不作为语义抽样依据"),
        ("抽样原则", "25单聊+25群聊；第一批10、第二批20、第三批20；每例固定100条；45/50为多人或双方可见互动，单发送者边界仅5/50；按机械结构分层，内容盲抽样"),
    ]
    ws["A4"] = "证据项"
    ws["B4"] = "值"
    style_header(ws, 4, 1, 2)
    for row_index, (label, value) in enumerate(audit_rows, start=5):
        ws.cell(row_index, 1, label)
        ws.cell(row_index, 2, value)
    style_body(ws, 5, 14, 1, 2, COLORS["gray2"])

    ws["A16"] = "批次"
    ws["B16"] = "单聊"
    ws["C16"] = "群聊"
    ws["D16"] = "合计"
    style_header(ws, 16, 1, 4)
    for row_index, batch in enumerate((1, 2, 3), start=17):
        direct = sum(1 for x in selected if x["batch"] == batch and x["clustertype"] == 1)
        group = sum(1 for x in selected if x["batch"] == batch and x["clustertype"] == 0)
        ws.append([])
        ws.cell(row_index, 1, f"第{batch}批")
        ws.cell(row_index, 2, direct)
        ws.cell(row_index, 3, group)
        ws.cell(row_index, 4, direct + group)
    style_body(ws, 17, 19, 1, 4, COLORS["blue"])

    ws["A22"] = "聊天类型"
    ws["B22"] = "发送者结构"
    ws["C22"] = "案例数"
    style_header(ws, 22, 1, 3)
    counter = Counter((x["chat_type_label"], sender_bucket(x["sender_count"])) for x in selected)
    row_index = 23
    for (chat_type, bucket), count in sorted(counter.items()):
        ws.cell(row_index, 1, chat_type)
        ws.cell(row_index, 2, bucket)
        ws.cell(row_index, 3, count)
        row_index += 1
    style_body(ws, 23, row_index - 1, 1, 3, COLORS["green"])
    ws.sheet_view.showGridLines = False
    set_widths(ws, {"A": 28, "B": 90, "C": 18, "D": 18, "E": 18, "F": 18, "G": 18, "H": 18})

    # 06 Original field index
    ws = wb.create_sheet("06_原表字段索引")
    style_title(ws, "原始 Excel 字段索引", "字段名称逐字来自原始 Excel 表头；本页不继承旧分析报告中的业务定义。", 4)
    ws.append([])
    ws.append(["原表列序号", "Excel列", "原表字段名与原注释", "本包用途"])
    style_header(ws, 4, 1, 4)
    used_fields = {
        2: "角色背景，只读",
        3: "群规模机械分层，只读",
        4: "消息溯源，只读",
        5: "案例/会话溯源，只读",
        6: "区分单聊与群聊，只读",
        8: "消息溯源，只读",
        10: "原始载荷保留，只读",
        11: "回复消息溯源，只读",
        12: "发送者溯源，只读",
        17: "原表顺序字段，只读",
        18: "主要消息正文，只读",
        19: "发言人显示名，只读",
        20: "可读时间，只读",
    }
    for index, field in enumerate(source_header, start=1):
        ws.append([index, get_column_letter(index), field, used_fields.get(index, "保留在原始输入；当前阅读表未展开")])
    style_body(ws, 5, ws.max_row, 1, 4, COLORS["gray2"])
    ws.freeze_panes = "A5"
    ws.auto_filter.ref = f"A4:D{ws.max_row}"
    ws.sheet_view.showGridLines = False
    set_widths(ws, {"A": 14, "B": 10, "C": 82, "D": 38})

    # Workbook-wide print and active sheet settings.
    wb.active = wb.sheetnames.index("03_Topic人工标注")
    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)

    case_manifest = []
    for item in selected:
        case_manifest.append({
            "batch": item["batch"],
            "batch_label": item["batch_label"],
            "case_id": item["case_id"],
            "chat_type": item["chat_type_label"],
            "clustertype": item["clustertype"],
            "clusterid": item["clusterid"],
            "message_count": item["message_count"],
            "sender_count": item["sender_count"],
            "sender_bucket": sender_bucket(item["sender_count"]),
            "role_pattern": item["role_pattern"],
            "group_size": item["group_size"],
            "first_time": item["first_time"].isoformat(sep=" ") if item["first_time"] else None,
            "last_time": item["last_time"].isoformat(sep=" ") if item["last_time"] else None,
            "span_days": round(item["span_days"], 4),
            "text_ratio": round(item["text_ratio"], 6),
            "structural_stratum": item["stratum"],
            "first_excel_row": item["first_excel_row"],
            "last_excel_row": item["last_excel_row"],
            "case_evidence_sha256": item["case_sha256"],
            "workbook_message_start_row": message_start_rows[item["case_id"]],
            "workbook_message_end_row": message_end_rows[item["case_id"]],
        })
    return {
        "cases": case_manifest,
        "message_rows": sum(len(selected_rows[(x["clustertype"], x["clusterid"])]) for x in selected),
    }


def main() -> None:
    if len(sys.argv) != 3:
        raise SystemExit("Usage: build_topic_pilot50_v1.py <original.xlsx> <output-dir>")
    source_path = Path(sys.argv[1]).resolve()
    output_dir = Path(sys.argv[2]).resolve()
    output_dir.mkdir(parents=True, exist_ok=True)

    source_sha = sha256(source_path)
    source_header, clusters, source_rows = scan_source(source_path)
    if source_rows != EXPECTED_SOURCE_ROWS:
        raise RuntimeError(f"Source row count changed: {source_rows} != {EXPECTED_SOURCE_ROWS}")
    if any(x["message_count"] != EXPECTED_MESSAGES_PER_CASE for x in clusters.values()):
        raise RuntimeError("Source is not uniformly grouped into 100-message cases")

    selected = choose_cases(clusters)
    if len(selected) != 50 or len({x["clusterid"] for x in selected}) != 50:
        raise RuntimeError("Selection must contain 50 unique clusterids")
    if Counter(x["clustertype"] for x in selected) != Counter({1: 25, 0: 25}):
        raise RuntimeError("Selection must contain 25 direct and 25 group cases")
    if Counter(x["batch"] for x in selected) != Counter({1: 10, 2: 20, 3: 20}):
        raise RuntimeError("Batch distribution must be 10/20/20")

    selected_rows = extract_selected_rows(source_path, selected)
    output_path = output_dir / "ClassIn_IM_Topic_Pilot50_人工Topic标注_v1.xlsx"
    result = build_workbook(
        source_path, source_sha, source_header, source_rows, clusters, selected, selected_rows, output_path
    )

    manifest = {
        "status": "TOPIC_PILOT50_V1_GENERATED_PENDING_QA",
        "version": VERSION,
        "generated_at": datetime.now().isoformat(timespec="seconds"),
        "research_boundary": "human topic discovery and method calibration; no prevalence or product conclusion",
        "source": {
            "path": str(source_path),
            "sha256": source_sha,
            "sheet": SOURCE_SHEET,
            "message_rows": source_rows,
            "case_pool": len(clusters),
            "direct_case_pool": sum(1 for x in clusters.values() if x["clustertype"] == 1),
            "group_case_pool": sum(1 for x in clusters.values() if x["clustertype"] == 0),
            "messages_per_case": EXPECTED_MESSAGES_PER_CASE,
        },
        "sampling": {
            "seed": SEED,
            "content_blind": True,
            "semantic_inputs": [],
            "prior_human_exposure_exclusion": {
                "clusterids": sorted(EXCLUDED_CLUSTER_IDS),
                "reason": "previously viewed in retired per-message trial; excluded to protect blind validation",
            },
            "structural_dimensions": [
                "clustertype", "visible sender count", "visible role pattern", "group-size field", "text completeness", "visible time span"
            ],
            "batch_distribution": {"batch_1": 10, "batch_2": 20, "batch_3": 20},
            "chat_type_distribution": {"direct": 25, "group": 25},
            "sender_bucket_quotas": {
                "direct": {"batch_1": {"2": 5}, "batch_2": {"1": 2, "2": 8}, "batch_3": {"1": 1, "2": 9}},
                "group": {"batch_1": {"3+ clear cases": 5}, "batch_2": {"1": 1, "2": 2, "3-5": 3, "6+": 4}, "batch_3": {"1": 1, "2": 2, "3-5": 3, "6+": 4}},
            },
        },
        "output": {
            "path": str(output_path),
            "sha256": sha256(output_path),
            "sheet_count": 7,
            "selected_cases": 50,
            "selected_message_rows": result["message_rows"],
            "human_input_sheet": "03_Topic人工标注",
            "human_input_columns": ["J", "K", "L", "M"],
        },
        "builder": {
            "path": str(Path(__file__).resolve()),
            "sha256": sha256(Path(__file__).resolve()),
        },
        "cases": result["cases"],
    }
    manifest_path = output_dir / "topic_pilot50_manifest_v1.json"
    manifest_path.write_text(json.dumps(manifest, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    manifest["manifest_path"] = str(manifest_path)
    print(json.dumps(manifest, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
