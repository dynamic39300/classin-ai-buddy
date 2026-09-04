#!/usr/bin/env python3
"""Build a deterministic, content-free Pilot-0 sample manifest.

The public manifest contains no raw cluster IDs, user IDs, message IDs, names,
or message text. A separate restricted map with cluster IDs is written outside
the repository for later controlled extraction.
"""

from __future__ import annotations

import argparse
import hashlib
import json
import math
import re
import sys
import time
from collections import Counter, defaultdict
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


EXPECTED_SHA256 = "a28f4c3125326c2e0f9086f8c0f67104671b94b2992a0173f4c34f90690fee6c"
DEFAULT_SEED = "classin-im-pilot0-20260829-v1"

QUOTAS = {
    "direct_two_senders": 6,
    "direct_one_sender": 2,
    "group_learner_and_staff": 8,
    "group_learner_only": 4,
    "group_staff_only": 4,
}


def file_sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def field_key(header: Any) -> str:
    text = "" if header is None else str(header).strip()
    return re.split(r"[（(]", text, maxsplit=1)[0].strip().lower()


def normalize(value: Any) -> str:
    if value is None or value == "":
        return "<EMPTY>"
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def is_blank(value: Any) -> bool:
    return value is None or value == "" or (isinstance(value, str) and not value.strip())


def is_empty_reference(value: Any) -> bool:
    if is_blank(value):
        return True
    if isinstance(value, (int, float)) and value == 0:
        return True
    return str(value).strip().lower() in {"0", "none", "null", "nan"}


def parse_epoch(value: Any) -> float | None:
    if value is None or value == "" or isinstance(value, bool):
        return None
    try:
        result = float(value)
    except (TypeError, ValueError):
        return None
    return result if math.isfinite(result) else None


def stable_score(seed: str, *parts: Any) -> str:
    text = "|".join([seed, *(str(part) for part in parts)])
    return hashlib.sha256(text.encode("utf-8")).hexdigest()


def sender_bucket(count: int) -> str:
    if count <= 1:
        return "1"
    if count == 2:
        return "2"
    if count <= 5:
        return "3-5"
    return "6+"


def time_span_bucket(min_epoch: float | None, max_epoch: float | None) -> str:
    if min_epoch is None or max_epoch is None:
        return "unknown"
    seconds = max_epoch - min_epoch
    if seconds <= 86400:
        return "<=1d"
    if seconds <= 7 * 86400:
        return "1d-7d"
    return "7d-27d"


@dataclass
class ClusterProfile:
    cluster_id: Any
    rows: int = 0
    channel_values: set[str] = field(default_factory=set)
    user_types: set[str] = field(default_factory=set)
    user_num_values: set[str] = field(default_factory=set)
    senders: set[Any] = field(default_factory=set)
    min_epoch: float | None = None
    max_epoch: float | None = None
    long_text_rows: int = 0
    missing_content_rows: int = 0
    non_object_msgdata_rows: int = 0
    explicit_reply_rows: int = 0

    def update_epoch(self, value: float | None) -> None:
        if value is None:
            return
        self.min_epoch = value if self.min_epoch is None else min(self.min_epoch, value)
        self.max_epoch = value if self.max_epoch is None else max(self.max_epoch, value)


def sample_group(profile: ClusterProfile) -> str | None:
    channel = next(iter(profile.channel_values)) if len(profile.channel_values) == 1 else "unknown"
    sender_count = len(profile.senders)
    if channel == "direct_1v1":
        if sender_count == 1:
            return "direct_one_sender"
        if sender_count >= 2:
            return "direct_two_senders"
        return None

    if channel != "class_group_declared":
        return None

    learner_roles = {"学生", "旁听生"}
    staff_roles = {"教师", "班主任"}
    has_learner = bool(profile.user_types & learner_roles)
    has_staff = bool(profile.user_types & staff_roles)
    if has_learner and has_staff:
        return "group_learner_and_staff"
    if has_learner and not has_staff:
        return "group_learner_only"
    if has_staff and profile.user_types <= staff_roles:
        return "group_staff_only"
    return None


def feature_key(profile: ClusterProfile) -> tuple[str, ...]:
    channel = next(iter(profile.channel_values)) if len(profile.channel_values) == 1 else "unknown"
    user_num = (
        next(iter(profile.user_num_values))
        if channel == "class_group_declared" and len(profile.user_num_values) == 1
        else "not_applicable_or_mixed"
    )
    structural_flags = []
    if profile.long_text_rows:
        structural_flags.append("long_text")
    if profile.missing_content_rows:
        structural_flags.append("missing_content")
    if profile.non_object_msgdata_rows:
        structural_flags.append("non_object_msgdata")
    if profile.explicit_reply_rows:
        structural_flags.append("explicit_reply")
    if not structural_flags:
        structural_flags.append("none")
    return (
        user_num,
        time_span_bucket(profile.min_epoch, profile.max_epoch),
        sender_bucket(len(profile.senders)),
        "+".join(structural_flags),
    )


def diverse_select(
    candidates: list[ClusterProfile], quota: int, seed: str, group_name: str
) -> list[ClusterProfile]:
    queues: dict[tuple[str, ...], list[ClusterProfile]] = defaultdict(list)
    for profile in candidates:
        queues[feature_key(profile)].append(profile)
    for key, queue in queues.items():
        queue.sort(key=lambda profile: stable_score(seed, group_name, key, profile.cluster_id))
    ordered_keys = sorted(queues, key=lambda key: stable_score(seed, group_name, "feature", key))

    selected: list[ClusterProfile] = []
    while len(selected) < quota:
        progressed = False
        for key in ordered_keys:
            if queues[key] and len(selected) < quota:
                selected.append(queues[key].pop(0))
                progressed = True
        if not progressed:
            break
    if len(selected) != quota:
        raise RuntimeError(
            f"Unable to satisfy quota for {group_name}: wanted {quota}, got {len(selected)}"
        )
    return selected


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("input_xlsx", type=Path)
    parser.add_argument("--public-output", type=Path, required=True)
    parser.add_argument("--restricted-map-output", type=Path, required=True)
    parser.add_argument("--seed", default=DEFAULT_SEED)
    parser.add_argument("--expected-sha256", default=EXPECTED_SHA256)
    parser.add_argument("--progress-every", type=int, default=100_000)
    args = parser.parse_args()

    started = time.time()
    actual_sha256 = file_sha256(args.input_xlsx)
    if actual_sha256 != args.expected_sha256:
        raise SystemExit(
            f"Input fingerprint mismatch: expected {args.expected_sha256}, got {actual_sha256}"
        )

    workbook = load_workbook(args.input_xlsx, read_only=True, data_only=False)
    worksheet = workbook.worksheets[0]
    rows = worksheet.iter_rows(values_only=True)
    headers = next(rows)
    keys = [field_key(value) for value in headers]
    index = {key: position for position, key in enumerate(keys)}

    profiles: dict[Any, ClusterProfile] = {}
    total_rows = 0
    for row in rows:
        total_rows += 1
        cluster_id = row[index["clusterid"]]
        profile = profiles.setdefault(cluster_id, ClusterProfile(cluster_id=cluster_id))
        profile.rows += 1
        clustertype = normalize(row[index["clustertype"]])
        profile.channel_values.add(
            "direct_1v1" if clustertype == "1" else "class_group_declared" if clustertype == "0" else "unknown"
        )
        profile.user_types.add(normalize(row[index["user_type"]]))
        profile.user_num_values.add(normalize(row[index["user_num"]]))
        sender = row[index["sourceuid"]]
        if not is_blank(sender):
            profile.senders.add(sender)
        profile.update_epoch(parse_epoch(row[index["timeformat"]]))

        content = row[index["concent"]]
        content_length = len(str(content)) if content is not None else 0
        if content_length > 1000:
            profile.long_text_rows += 1
        if is_blank(content):
            profile.missing_content_rows += 1

        msgdata = row[index["msgdata"]]
        msgdata_text = str(msgdata).strip() if msgdata is not None else ""
        if not (msgdata_text.startswith("{") and msgdata_text.endswith("}")):
            profile.non_object_msgdata_rows += 1

        if not is_empty_reference(row[index["replymsgid"]]):
            profile.explicit_reply_rows += 1

        if args.progress_every and total_rows % args.progress_every == 0:
            print(f"scanned_rows={total_rows}", file=sys.stderr, flush=True)

    candidates_by_group: dict[str, list[ClusterProfile]] = defaultdict(list)
    excluded_profiles = 0
    invalid_window_profiles = 0
    for profile in profiles.values():
        if profile.rows != 100:
            invalid_window_profiles += 1
            continue
        group = sample_group(profile)
        if group is None:
            excluded_profiles += 1
        else:
            candidates_by_group[group].append(profile)

    selected_records: list[dict[str, Any]] = []
    restricted_records: list[dict[str, Any]] = []
    candidate_counts = {group: len(candidates_by_group[group]) for group in QUOTAS}

    group_prefix = {
        "direct_two_senders": "DIR2",
        "direct_one_sender": "DIR1",
        "group_learner_and_staff": "GMIX",
        "group_learner_only": "GLEARN",
        "group_staff_only": "GSTAFF",
    }
    segment_use = Counter()

    for group_name, quota in QUOTAS.items():
        selected = diverse_select(candidates_by_group[group_name], quota, args.seed, group_name)
        for ordinal, profile in enumerate(selected, 1):
            sample_id = f"P0-{group_prefix[group_name]}-{ordinal:02d}"
            preferred_segments = sorted(
                range(5),
                key=lambda segment: stable_score(
                    args.seed, "segment-preference-v1-1", profile.cluster_id, segment
                ),
            )
            minimum_use = min(segment_use.get(segment, 0) for segment in range(5))
            segment_index = next(
                segment
                for segment in preferred_segments
                if segment_use.get(segment, 0) == minimum_use
            )
            segment_use[segment_index] += 1
            core_start = segment_index * 20 + 1
            core_end = core_start + 19
            context_start = max(1, core_start - 5)
            context_end = min(100, core_end + 5)
            channel = next(iter(profile.channel_values))
            role_composition = " | ".join(sorted(profile.user_types))
            user_num = (
                next(iter(profile.user_num_values))
                if channel == "class_group_declared" and len(profile.user_num_values) == 1
                else "not_applicable_direct"
            )
            flags = []
            if profile.long_text_rows:
                flags.append("contains_long_text")
            if profile.missing_content_rows:
                flags.append("contains_missing_content")
            if profile.non_object_msgdata_rows:
                flags.append("contains_non_object_msgdata")
            if profile.explicit_reply_rows:
                flags.append("contains_explicit_reply")
            if not flags:
                flags.append("no_selected_structural_flag")

            public_record = {
                "sample_id": sample_id,
                "sample_group": group_name,
                "channel": channel,
                "observed_role_composition": role_composition,
                "user_num_bucket": user_num,
                "observed_sender_count_bucket": sender_bucket(len(profile.senders)),
                "window_time_span_bucket": time_span_bucket(profile.min_epoch, profile.max_epoch),
                "structural_flags": flags,
                "core_rn_start": core_start,
                "core_rn_end": core_end,
                "context_rn_start": context_start,
                "context_rn_end": context_end,
                "selection_status": "selected",
                "replacement_for": None,
                "loss_reason": None,
                "estimation_eligibility": "none_method_validation_only",
            }
            selected_records.append(public_record)
            restricted_records.append(
                {
                    "sample_id": sample_id,
                    "clusterid": profile.cluster_id,
                    "core_rn_start": core_start,
                    "core_rn_end": core_end,
                    "context_rn_start": context_start,
                    "context_rn_end": context_end,
                }
            )

    selected_records.sort(key=lambda record: record["sample_id"])
    restricted_records.sort(key=lambda record: record["sample_id"])
    segment_distribution = Counter(
        f"{record['core_rn_start']}-{record['core_rn_end']}" for record in selected_records
    )

    public_manifest = {
        "manifest_id": "classin-im-pilot0-public-manifest-v1-1-20260829",
        "protocol_version": "P2-v0.1",
        "selection_method_version": "v1.1",
        "source": {
            "file_name": args.input_xlsx.name,
            "sha256": actual_sha256,
            "export_window_count": len(profiles),
            "rows_scanned": total_rows,
        },
        "selection": {
            "seed": args.seed,
            "quota": QUOTAS,
            "candidate_counts": candidate_counts,
            "selected_count": len(selected_records),
            "invalid_window_profiles": invalid_window_profiles,
            "profiles_outside_pilot_groups": excluded_profiles,
            "core_segment_distribution": dict(segment_distribution.most_common()),
            "method": "balanced purposive coverage with deterministic structural-diversity round-robin",
            "estimation_warning": "Pilot-0 is not a probability sample and must not be used for prevalence estimates.",
        },
        "samples": selected_records,
        "privacy": "No raw cluster IDs, user IDs, message IDs, names, or message text are present.",
    }

    restricted_map = {
        "classification": "RESTRICTED_RESEARCH_MAPPING_DO_NOT_COMMIT",
        "source_path": str(args.input_xlsx),
        "source_sha256": actual_sha256,
        "seed": args.seed,
        "protocol_version": "P2-v0.1",
        "selection_method_version": "v1.1",
        "samples": restricted_records,
    }

    args.public_output.parent.mkdir(parents=True, exist_ok=True)
    args.restricted_map_output.parent.mkdir(parents=True, exist_ok=True)
    args.public_output.write_text(
        json.dumps(public_manifest, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    args.restricted_map_output.write_text(
        json.dumps(restricted_map, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    workbook.close()

    print(f"public_output={args.public_output}")
    print(f"restricted_map_output={args.restricted_map_output}")
    print(f"selected_count={len(selected_records)}")
    print(f"elapsed_seconds={round(time.time() - started, 2)}")


if __name__ == "__main__":
    main()
