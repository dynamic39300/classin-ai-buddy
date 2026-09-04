#!/usr/bin/env python3
"""Validate Topic Pilot 50 workbook against the locked original Excel."""

from __future__ import annotations

import hashlib
import json
import sys
from collections import Counter, defaultdict
from datetime import datetime
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


EXPECTED_SHEETS = [
    "00_使用说明",
    "01_案例目录",
    "02_案例消息",
    "03_Topic人工标注",
    "04_问题与校准",
    "05_抽样审计",
    "06_原表字段索引",
]
EXCLUDED_CLUSTER_IDS = {"27324211"}


def text(value: Any) -> str:
    return "" if value is None else str(value)


def sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def main() -> None:
    if len(sys.argv) != 5:
        raise SystemExit("Usage: validate_topic_pilot50_v1.py <source.xlsx> <pilot.xlsx> <manifest.json> <qa.json>")
    source_path, pilot_path, manifest_path, qa_path = map(lambda p: Path(p).resolve(), sys.argv[1:])
    manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
    checks: dict[str, Any] = {}

    wb = load_workbook(pilot_path, data_only=False)
    checks["sheet_names_exact"] = wb.sheetnames == EXPECTED_SHEETS
    checks["active_sheet"] = wb.active.title

    topic_ws = wb["03_Topic人工标注"]
    topic_rows = list(topic_ws.iter_rows(min_row=5, max_row=54, values_only=False))
    checks["topic_case_rows"] = len(topic_rows)
    checks["human_topic_cells_blank"] = all(row[col - 1].value in (None, "") for row in topic_rows for col in (10, 11, 12))
    checks["status_defaults"] = Counter(text(row[12].value) for row in topic_rows)
    checks["status_validation_ranges"] = [str(x.sqref) for x in topic_ws.data_validations.dataValidation]
    checks["topic_link_formulas"] = all(text(row[8].value).startswith("=HYPERLINK(") for row in topic_rows)

    msg_ws = wb["02_案例消息"]
    checks["message_rows"] = msg_ws.max_row - 4
    checks["audit_columns_hidden"] = all(msg_ws.column_dimensions[col].hidden for col in "JKLMNOPQRST")
    counts: Counter[str] = Counter()
    workbook_rows: dict[int, tuple[Any, ...]] = {}
    case_seq: defaultdict[str, list[int]] = defaultdict(list)
    selected_clusterids: set[str] = set()
    for row in msg_ws.iter_rows(min_row=5, values_only=True):
        case_id = text(row[0])
        counts[case_id] += 1
        case_seq[case_id].append(int(row[3]))
        excel_row = int(row[9])
        if excel_row in workbook_rows:
            raise RuntimeError(f"Duplicate original Excel row in pilot: {excel_row}")
        workbook_rows[excel_row] = row
        selected_clusterids.add(text(row[10]))
    checks["case_count"] = len(counts)
    checks["messages_per_case"] = Counter(counts.values())
    checks["case_seq_exact"] = all(seqs == list(range(1, 101)) for seqs in case_seq.values())
    checks["excluded_cluster_absent"] = not (selected_clusterids & EXCLUDED_CLUSTER_IDS)

    catalog_ws = wb["01_案例目录"]
    catalog = list(catalog_ws.iter_rows(min_row=5, max_row=54, values_only=True))
    checks["catalog_cases"] = len(catalog)
    checks["batch_distribution"] = Counter(text(row[0]).split("｜")[0] for row in catalog)
    checks["chat_type_distribution"] = Counter(text(row[2]) for row in catalog)
    checks["sender_buckets_by_type"] = Counter(f"{text(row[2])}|{text(row[6])}" for row in catalog)
    checks["single_sender_boundary_cases"] = sum(1 for row in catalog if text(row[6]) == "1")
    checks["multi_or_two_sender_cases"] = sum(1 for row in catalog if text(row[6]) != "1")
    checks["direct_single_sender_cases"] = sum(1 for row in catalog if text(row[2]) == "1v1单聊" and text(row[6]) == "1")
    checks["group_single_sender_cases"] = sum(1 for row in catalog if text(row[2]) == "班级群聊" and text(row[6]) == "1")
    checks["group_two_sender_cases"] = sum(1 for row in catalog if text(row[2]) == "班级群聊" and text(row[6]) == "2")
    checks["group_multi_sender_cases"] = sum(1 for row in catalog if text(row[2]) == "班级群聊" and text(row[6]) in {"3-5", "6+"})
    checks["catalog_links"] = all(text(row[16]).startswith("=HYPERLINK(") for row in catalog)

    # Compare every exported evidence field with the exact original Excel row.
    source_wb = load_workbook(source_path, read_only=True, data_only=True)
    source_ws = source_wb["Sheet"]
    wanted = set(workbook_rows)
    matched = 0
    mismatches: list[dict[str, Any]] = []
    for excel_row, source_row in enumerate(source_ws.iter_rows(min_row=2, values_only=True), start=2):
        if excel_row not in wanted:
            continue
        pilot = workbook_rows[excel_row]
        pairs = {
            "time": (pilot[4], source_row[19]),
            "strtalker": (pilot[5], source_row[18]),
            "user_type": (pilot[6], source_row[1]),
            "sourceuid": (pilot[7], source_row[11]),
            "content": (pilot[8], source_row[17]),
            "clusterid": (pilot[10], source_row[4]),
            "clustertype": (pilot[11], source_row[5]),
            "id": (pilot[12], source_row[3]),
            "msgid": (pilot[13], source_row[7]),
            "replymsgid": (pilot[14], source_row[10]),
            "rn": (pilot[15], source_row[16]),
            "identity": (pilot[16], source_row[0]),
            "user_num": (pilot[17], source_row[2]),
            "targetuids": (pilot[18], source_row[12]),
            "msgdata": (pilot[19], source_row[9]),
        }
        bad = [name for name, (left, right) in pairs.items() if text(left) != text(right)]
        if bad:
            mismatches.append({"excel_row": excel_row, "fields": bad})
        matched += 1
    source_wb.close()
    wb.close()
    checks["original_rows_matched"] = matched
    checks["evidence_mismatch_count"] = len(mismatches)
    checks["evidence_mismatch_examples"] = mismatches[:10]

    expected = {
        "sheet_names_exact": True,
        "active_sheet": "03_Topic人工标注",
        "topic_case_rows": 50,
        "human_topic_cells_blank": True,
        "status_defaults": Counter({"未开始": 50}),
        "status_validation_ranges": ["M5:M54"],
        "topic_link_formulas": True,
        "message_rows": 5000,
        "audit_columns_hidden": True,
        "case_count": 50,
        "messages_per_case": Counter({100: 50}),
        "case_seq_exact": True,
        "excluded_cluster_absent": True,
        "catalog_cases": 50,
        "batch_distribution": Counter({"第一批": 10, "第二批": 20, "第三批": 20}),
        "chat_type_distribution": Counter({"1v1单聊": 25, "班级群聊": 25}),
        "single_sender_boundary_cases": 5,
        "multi_or_two_sender_cases": 45,
        "direct_single_sender_cases": 3,
        "group_single_sender_cases": 2,
        "group_two_sender_cases": 4,
        "group_multi_sender_cases": 19,
        "catalog_links": True,
        "original_rows_matched": 5000,
        "evidence_mismatch_count": 0,
    }
    failures = {key: {"actual": checks.get(key), "expected": value} for key, value in expected.items() if checks.get(key) != value}
    report = {
        "status": "PASS" if not failures else "FAIL",
        "validated_at": datetime.now().isoformat(timespec="seconds"),
        "source": {"path": str(source_path), "sha256": sha256(source_path)},
        "workbook": {"path": str(pilot_path), "sha256": sha256(pilot_path)},
        "manifest": {"path": str(manifest_path), "declared_workbook_sha256": manifest["output"]["sha256"]},
        "checks": checks,
        "failures": failures,
    }
    qa_path.write_text(json.dumps(report, ensure_ascii=False, indent=2, default=lambda obj: dict(obj)) + "\n", encoding="utf-8")
    print(json.dumps(report, ensure_ascii=False, indent=2, default=lambda obj: dict(obj)))
    if failures:
        raise SystemExit(1)


if __name__ == "__main__":
    main()
