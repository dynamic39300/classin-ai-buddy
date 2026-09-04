#!/usr/bin/env python3
"""Create Topic Pilot 50 v1.1 with one human-input column.

The migration preserves all evidence and existing human Topic text from v1.
It removes the three unused human fields and records the five-effective-
utterance Topic threshold in the workbook instructions.
"""

from __future__ import annotations

import hashlib
import json
import sys
from datetime import datetime
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side


NAVY = "17384B"
WHITE = "FFFFFF"
GRAY2 = "F4F6F7"
YELLOW = "FFF2CC"
BORDER_COLOR = "B8C4CC"
THIN = Side(style="thin", color=BORDER_COLOR)
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def value_hash(values: list[Any]) -> str:
    payload = "\n\x1e\n".join("" if value is None else str(value) for value in values)
    return hashlib.sha256(payload.encode("utf-8")).hexdigest()


def style_instruction_row(ws: Any, row: int, label: str, explanation: str) -> None:
    ws.cell(row, 1, label)
    ws.cell(row, 1).fill = PatternFill("solid", fgColor=NAVY)
    ws.cell(row, 1).font = Font(color=WHITE, bold=True)
    ws.cell(row, 1).alignment = Alignment(vertical="top", wrap_text=True)
    if f"B{row}:J{row}" not in {str(item) for item in ws.merged_cells.ranges}:
        ws.merge_cells(start_row=row, start_column=2, end_row=row, end_column=10)
    ws.cell(row, 2, explanation)
    ws.cell(row, 2).fill = PatternFill("solid", fgColor=GRAY2)
    ws.cell(row, 2).alignment = Alignment(vertical="top", wrap_text=True)
    for col in range(1, 11):
        ws.cell(row, col).border = BORDER
    ws.row_dimensions[row].height = 48


def main() -> None:
    if len(sys.argv) != 4:
        raise SystemExit("Usage: simplify_topic_pilot50_v1_1.py <v1.xlsx> <v1.1.xlsx> <manifest.json>")
    source_path, output_path, manifest_path = map(lambda item: Path(item).resolve(), sys.argv[1:])
    wb = load_workbook(source_path)
    if wb.sheetnames != [
        "00_使用说明", "01_案例目录", "02_案例消息", "03_Topic人工标注",
        "04_问题与校准", "05_抽样审计", "06_原表字段索引",
    ]:
        raise RuntimeError("Unexpected v1 workbook structure")

    topic_ws = wb["03_Topic人工标注"]
    original_topics = [topic_ws.cell(row, 10).value for row in range(5, 55)]
    if sum(bool(value) for value in original_topics[:10]) != 10:
        raise RuntimeError("Expected all 10 Batch-1 Topic cells to be filled")
    if any(value not in (None, "") for value in original_topics[10:]):
        raise RuntimeError("Expected Batch-2/3 Topic cells to remain blank")
    if any(topic_ws.cell(row, col).value not in (None, "") for row in range(5, 55) for col in (11, 12)):
        raise RuntimeError("K/L contain unexpected human input")

    # Keep J as the sole human-input column and physically remove K:M.
    # Unmerge title bands first because deleting merged columns can leave
    # stale merged-cell coordinates in openpyxl.
    for merged in ("A1:M1", "A2:M2"):
        if merged in {str(item) for item in topic_ws.merged_cells.ranges}:
            topic_ws.unmerge_cells(merged)
    topic_ws.delete_cols(11, 3)
    topic_ws.merge_cells("A1:J1")
    topic_ws.merge_cells("A2:J2")
    topic_ws["A1"] = "Topic Pilot 50 v1.1 · 人工 Topic 初稿"
    topic_ws["A2"] = "只填写黄色 J 列。每个案例一行；同一主题至少有5句有效发言才进入 Topic，少量插话不要求记录。"
    topic_ws["J4"] = "人工Topic初稿【唯一填写列；一行一个Topic；同一主题≥5句有效发言】"
    topic_ws.auto_filter.ref = "A4:J54"
    topic_ws.freeze_panes = "J5"
    topic_ws.data_validations.dataValidation = []
    topic_ws.column_dimensions["J"].width = 72
    for row in range(5, 55):
        topic_ws.cell(row, 10).fill = PatternFill("solid", fgColor=YELLOW)
        topic_ws.cell(row, 10).alignment = Alignment(vertical="top", wrap_text=True)

    # Update the instruction sheet without adding another manual task.
    guide_ws = wb["00_使用说明"]
    guide_ws["A1"] = "ClassIn IM Topic Pilot 50 v1.1 · 仅填写 Topic"
    guide_ws["A2"] = "人工以100条消息案例为单位提炼主要 Topic；单条消息只作证据。人工只填写03表黄色J列。"
    style_instruction_row(
        guide_ws,
        4,
        "你只需要填写哪里",
        "只填写 03_Topic人工标注 的黄色 J 列“人工Topic初稿”。不再填写关键证据、上下文不足或标注状态。",
    )
    style_instruction_row(
        guide_ws,
        7,
        "Topic进入门槛",
        "同一语义主题至少出现5句有效发言才写入 Topic；少于5句的零散表达、短暂插话或独立问题不要求记录，也不算漏标。",
    )
    style_instruction_row(
        guide_ws,
        8,
        "什么是有效发言",
        "按有语义的讨论发言计数，不机械按Excel行计数。纯表情、重复发送、无语义符号，以及同一人连续拆开的句子片段不能虚增权重。",
    )
    style_instruction_row(
        guide_ws,
        12,
        "发言数量的含义",
        "数量衡量 Topic 在聊天中的显著性与关注度，不等于问题严重程度、业务价值或产品优先级；短消息高影响事件不混入当前 Topic 口径。",
    )

    audit_ws = wb["05_抽样审计"]
    audit_ws["A15"] = "Topic收录门槛"
    audit_ws["B15"] = "同一主题≥5句有效发言；纯表情、重复、无语义符号和拆句不机械增权"
    for col in (1, 2):
        audit_ws.cell(15, col).fill = PatternFill("solid", fgColor=GRAY2)
        audit_ws.cell(15, col).alignment = Alignment(vertical="top", wrap_text=True)
        audit_ws.cell(15, col).border = BORDER

    wb.active = wb.sheetnames.index("03_Topic人工标注")
    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)

    manifest = {
        "status": "TOPIC_PILOT50_V1_1_READY_FOR_BATCH2_HUMAN_TOPIC_ANNOTATION",
        "version": "v1.1",
        "generated_at": datetime.now().isoformat(timespec="seconds"),
        "source_workbook": {"path": str(source_path), "sha256": sha256(source_path)},
        "output_workbook": {"path": str(output_path), "sha256": sha256(output_path)},
        "migration": {
            "sole_human_input": "03_Topic人工标注!J5:J54",
            "removed_human_fields": ["关键证据", "上下文不足/不确定说明", "标注状态"],
            "topic_admission_rule": "same semantic topic has at least 5 effective utterances",
            "effective_utterance_exclusions": ["pure emoji", "duplicates", "non-semantic symbols", "split fragments"],
            "utterance_count_meaning": "conversational salience, not severity, business value, or product priority",
        },
        "preservation": {
            "batch1_filled_topic_cells": 10,
            "batch2_3_blank_topic_cells": 40,
            "topic_values_sha256": value_hash(original_topics),
            "evidence_sheets_preserved": ["01_案例目录", "02_案例消息", "06_原表字段索引"],
        },
        "builder": {"path": str(Path(__file__).resolve()), "sha256": sha256(Path(__file__).resolve())},
    }
    manifest_path.write_text(json.dumps(manifest, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps(manifest, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
